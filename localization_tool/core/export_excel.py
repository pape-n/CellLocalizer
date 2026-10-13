"""Results as one Excel file (SPEC §4, D8, D31, D39).

Sheets, in this order:
- Read me: what is where, what the numbers mean
- Summary: one row per cell and target, columns grouped like the table in the program
  (Signal %, Volume %, Enrichment per region), sorted by target, condition, experiment and image
- one sheet per key measure and target in Prism layout (columns = conditions, rows = cells),
  ready to paste into a Prism column table, with the cell IDs in the same order to the right
- Replicates: mean per experiment and condition (cells of one coverslip are not independent)
Prism sheets and Replicates hold only checked cells (status accepted). If no cell is checked yet
(e.g. a command line run), they hold all cells and say so.
- Excluded: cells and images left out, with the reason (only if there are any)
- All values: every number per cell and target
- Profiles: distance and height profiles
- Acquisition: microscope settings per image and channel (D63)
- Info: program version, date, parameters, channel roles, images
"""

from __future__ import annotations

import json
import time
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import __version__

ACCENT = "2A6496"
LIGHT = "DCE6F0"
BAND = "F3F6F9"
PCT, ENR, VOL, R2 = "0.0", "0.00", "0", "0.00"

# key measures: (column in the cell rows, short sheet title, explanation)
KEY_MEASURES = [
    ("enrichment_nucleus", "Enr nucleus", "Enrichment in the nucleus (1 = as dense as the cell average)"),
    ("enrichment_cytoplasm", "Enr cytoplasm", "Enrichment in the cytoplasm (everything outside the nucleus)"),
    ("enrichment_actin", "Enr actin", "Enrichment in the actin part of the cytoplasm"),
    ("signal_nucleus_pct", "Sig nucleus", "Share of the signal in the nucleus, %"),
    ("signal_cytoplasm_pct", "Sig cytoplasm", "Share of the signal in the cytoplasm, % (nucleus + cytoplasm = 100)"),
    ("signal_actin_pct", "Sig actin", "Share of the signal in the actin part of the cytoplasm, %"),
    ("enrichment_actin_rich", "Enr actin-rich", "Enrichment in the actin-rich zone (front)"),
    ("enrichment_basal_layer", "Enr basal", "Enrichment in the basal layer above the coverslip"),
    ("enrichment_front", "Enr front", "Enrichment in the front half of the cell (axis from the actin channel)"),
]

# Summary: groups of (key, header, number format)
_SVE = lambda k: [(f"signal_{k}_pct", "Signal %", PCT), (f"volume_{k}_pct", "Volume %", PCT),  # noqa: E731
                  (f"enrichment_{k}", "Enrichment", ENR)]
SUMMARY_GROUPS = [
    ("Cell", [("experiment", "Experiment", None), ("condition", "Condition", None), ("image", "Image", None),
              ("cell", "Cell", None), ("target", "Target", None), ("status", "Status", None)]),
    ("Nucleus", _SVE("nucleus")),
    ("Cytoplasm", _SVE("cytoplasm")),
    ("Actin (part of cytoplasm)", _SVE("actin")),
    ("Actin-rich zone", _SVE("actin_rich")),
    ("Nucleus interior", _SVE("nucleus_interior")),
    ("Nuclear envelope zone", _SVE("envelope_zone")),
    ("Cortex shell", _SVE("cortex_shell")),
    ("Basal layer", _SVE("basal_layer")),
    ("Front half", [("signal_front_pct", "Signal %", PCT), ("volume_front_pct", "Volume %", PCT),
                    ("enrichment_front", "Enrichment", ENR)]),
    ("Target mask", [("target_mask_signal_pct", "Signal %", PCT), ("target_mask_volume_pct", "Volume %", PCT)]),
    ("With actin", [("pcc_actin", "PCC", R2), ("m1_actin", "M1", R2), ("m2_actin", "M2", R2)]),
    ("Sensitivity (edges x0.8 / x1.2)", [("enrichment_nucleus_min", "E nucleus min", ENR),
                                         ("enrichment_nucleus_max", "E nucleus max", ENR),
                                         ("signal_nucleus_pct_min", "Sig nucleus min", PCT),
                                         ("signal_nucleus_pct_max", "Sig nucleus max", PCT)]),
    ("Cell size and shape", [("cell_volume_um3", "Cell µm³", VOL), ("nucleus_volume_um3", "Nucleus µm³", VOL),
                             ("cell_height_um", "Height µm", PCT), ("polarity_index", "Polarity index", R2)]),
    ("", [("notes", "Notes", None)]),
]

STATUS_TEXT = {"automatic": "to check"}     # stored as "automatic", shown as "to check" like in the program


def _read_me(params: dict) -> list:
    env = params.get("envelope_um", 0.3)
    cortex = params.get("cortex_um", 0.5)
    basal = params.get("basal_um", 0.5)
    bin_um = params.get("profile_bin_um", 0.2)
    return [s if not callable(s) else s(env, cortex, basal, bin_um) for s in _READ_ME]


_READ_ME = [
    ("What is in this file", None),
    ("Summary", "One row per cell and target, grouped like the table in the program. Sorted by target, "
                "condition, experiment and image."),
    ("<target> Enr ... / Sig ...", "One sheet per key number and target in Prism layout: one column per condition, "
                                   "one row per cell. Select the block from the condition names down and paste it "
                                   "into a Prism column table. To the right, after an empty column, the image and "
                                   "cell of every value in the same order. Only checked cells (accepted)."),
    ("Replicates", "Mean per experiment and condition of the checked cells. Cells of one coverslip are not "
                   "independent; for statistics use one value per experiment (or show both as a SuperPlot)."),
    ("Excluded", "Cells and images you left out, with your reason. They are not in any other sheet."),
    ("All values", "Every number the program computed, one row per cell and target."),
    lambda e, c, b, bn: ("Profiles", "Signal and volume along the distance to the nucleus, the depth below the "
                                     f"cell surface and the height above the coverslip, in bins of {bn:g} µm."),
    ("Acquisition", "Microscope settings per image and channel from the files, and which differ within an experiment."),
    ("Info", "Program version, parameters, channel roles and images, so every number can be reproduced."),
    ("", None),
    ("What the numbers mean", None),
    ("Signal %", "Share of the target's total signal in the cell that lies in a region. "
                 "Nucleus + cytoplasm = 100 %."),
    ("Volume %", "Share of the cell volume that the region takes up."),
    ("Enrichment", "Signal % divided by Volume %. 1 = as dense as the cell average, above 1 enriched, "
                   "below 1 depleted. Best number to compare conditions, because it corrects for regions "
                   "that change size, e.g. when activated cells spread."),
    ("Cytoplasm and actin", "Cytoplasm is everything in the cell outside the nucleus, including the actin. "
                            "Actin is the part of the cytoplasm with actin signal; actin-rich is its brightest "
                            "part. A voxel can count for cytoplasm and actin at the same time."),
    lambda e, c, b, bn: ("Zones", f"Nucleus interior (more than {e:g} µm inside the nuclear border), nuclear "
                                  f"envelope zone (within {e:g} µm of it), cortex shell ({c:g} µm below the cell "
                                  f"surface), basal layer ({b:g} µm above the coverslip). Zones overlap, they are "
                                  "extra views."),
    ("Front half", "Axis in the coverslip plane from the cell centre to the actin-weighted centre. Polarity "
                   "index = axis length / radius of the cell footprint: about 0 for an even cortex, larger with "
                   "a front. For cells with a low index the front values mean little."),
    ("Target mask", "Voxels with clear target signal. A visual check and an extra view; the shares above use "
                    "all of the target signal and do not depend on this mask."),
    ("PCC, M1, M2", "Colocalization with actin inside the cell, for comparison with the literature. PCC: "
                    "Pearson correlation (-1 to 1). M1: share of the target signal where actin is above its "
                    "Costes threshold, M2 the other way round. Empty when the channels are not positively "
                    "correlated."),
    ("Sensitivity", "The main values computed again with the cell edge and the nucleus edge at 0.8 and 1.2 "
                    "times their threshold (four variants). Min and max include the value itself. A wide span "
                    "means the value depends on where exactly the edge lies: look at that cell again."),
    ("Acquisition", "Airyscan strength, laser transmission and detector gain per image and channel, read from "
                    "the CZI file. Within one experiment they should be the same; a note says when they differ. "
                    "Automatic Airyscan strength differs from image to image by design."),
    ("Lowest plane", "Column bottom_plane_ratio: brightness of the plane at the coverslip relative to the "
                     "brightest plane within the cell footprint (actin, targets). From 0.6 on the note 'lowest "
                     "plane already bright' appears: the stack probably starts inside the cell, so the basal "
                     "layer is the lowest recorded layer, not the layer on the coverslip. Start stacks about "
                     "1 µm below the cell."),
    ("Channel offset", "Column channel_offset_correction: the chromatic offset between channels measured with "
                       "beads and corrected in whole voxels, or 'none'. Without a correction, heights and zones "
                       "near surfaces (basal layer, cortex, nuclear envelope) can be shifted by the optics."),
    ("Status", "accepted: you checked the cell. to check: computed without your check (e.g. 'Analyze all "
               "images'); these cells are in Summary and All values, not in the Prism sheets and Replicates."),
    ("Empty cells", "No value: e.g. no nucleus or actin channel, no nucleus found, no Costes threshold, or "
                    "no target mask because the target channel shows no staining above noise. "
                    "The column Notes says why."),
]

THIN = Side(style="thin", color="B0BEC5")


def _header_cell(ws, row, col, value, strong=True):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=True, color="FFFFFF" if strong else "1E2227")
    c.fill = PatternFill("solid", fgColor=ACCENT if strong else LIGHT)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return c


def _header(ws, row: int, values: list[str]) -> None:
    for c, v in enumerate(values, start=1):
        _header_cell(ws, row, c, v)


def _widths(ws, widths: dict[int, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _num(v):
    if isinstance(v, float) and v != v:
        return None
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)   # e.g. a tab pasted into a note
    return v


def _checked(rows: list[dict]) -> tuple[list[dict], bool]:
    """Rows of checked cells for the Prism sheets and Replicates, and whether that selection was made."""
    ok = [r for r in rows if str(r.get("status", "")).startswith("accepted")]
    return (ok, True) if ok else (rows, False)


def _sheet_title(text: str, used: set[str]) -> str:
    bad = '[]:*?/\\'
    t = "".join("_" if ch in bad else ch for ch in text)[:31]
    base, k = t, 2
    while t in used:
        t = f"{base[:28]}_{k}"
        k += 1
    used.add(t)
    return t


def _sorted(rows: list[dict]) -> list[dict]:
    return sorted(rows, key=lambda r: (str(r.get("target", "")), str(r.get("condition", "")),
                                       str(r.get("experiment", "")), str(r.get("image", "")), str(r.get("cell", ""))))


ACQ_PREFIXES = ("airyscan_strength_", "laser_transmission_pct_", "detector_voltage_", "detector_offset_")
ACQ_NAMES = {"airyscan_strength_": "Airyscan strength", "laser_transmission_pct_": "laser %",
             "detector_voltage_": "detector gain (V)", "detector_offset_": "detector offset"}


def acquisition_notes(rows: list[dict], tolerance: float = 0.01) -> list[dict]:
    """Copies of the rows with a note where acquisition settings differ within an experiment (D63).
    Values are compared per image (one value per image and channel), relative tolerance 1 %."""
    by_exp: dict = {}
    for row in rows:
        by_exp.setdefault(row.get("experiment") or "", []).append(row)
    out = []
    for exp, grp in by_exp.items():
        texts = []
        cols = sorted({k for r in grp for k in r if k.startswith(ACQ_PREFIXES)})
        for col in cols:
            if col.startswith("airyscan_strength_") and all(
                    str(r.get("airyscan_mode_" + col[len("airyscan_strength_"):]) or "").lower().startswith("auto")
                    for r in grp if col in r):
                continue        # automatic strength is estimated per image and differs by design
            per_image = {}
            for r in grp:
                v = _num(r.get(col))
                if isinstance(v, (int, float)):
                    per_image[(r.get("condition"), r.get("image"))] = float(v)   # same file name in two folders
            if len(per_image) < 2:
                continue
            lo, hi = min(per_image.values()), max(per_image.values())
            if hi - lo > tolerance * max(abs(hi), 1e-9):
                prefix = next(p for p in ACQ_PREFIXES if col.startswith(p))
                texts.append(f"{ACQ_NAMES[prefix]} {col[len(prefix):]} {lo:g} to {hi:g}")
        note = ("acquisition differs within experiment: " + ", ".join(texts)) if texts else ""
        for r in grp:
            r = dict(r)
            if note:
                r["notes"] = "; ".join(t for t in (r.get("notes") or "", note) if t)
                r["_acquisition_differs"] = ", ".join(texts)
            out.append(r)
    return out


def _summary(ws, rows: list[dict]) -> None:
    col = 1
    for title, cols in SUMMARY_GROUPS:
        if len(cols) > 1:
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(cols) - 1)
        _header_cell(ws, 1, col, title)
        for k, (_, head, _) in enumerate(cols):
            _header_cell(ws, 2, col + k, head, strong=False)
        col += len(cols)
    flat = [c for _, cols in SUMMARY_GROUPS for c in cols]
    group_start = set()
    col = 1
    for _, cols in SUMMARY_GROUPS:
        group_start.add(col)
        col += len(cols)
    band, last = False, None
    for i, row in enumerate(_sorted(rows), start=3):
        key = (row.get("target"), row.get("condition"))
        if key != last:
            band, last = not band, key
        for c, (k, _, fmt) in enumerate(flat, start=1):
            v = row.get(k)
            cell = ws.cell(row=i, column=c, value=_num(STATUS_TEXT.get(v, v) if k == "status" else v))
            if fmt:
                cell.number_format = fmt
            if band:
                cell.fill = PatternFill("solid", fgColor=BAND)
            if c in group_start:
                cell.border = Border(left=THIN)
    ws.freeze_panes = "G3"
    ws.row_dimensions[2].height = 30
    _widths(ws, {c: 10 for c in range(1, len(flat) + 1)})
    _widths(ws, {1: 12, 2: 16, 3: 36, 4: 8, 5: 8, 6: 11, len(flat): 60})


def write_results(path: str | Path, rows: list[dict], profile_rows: list[dict], info: dict,
                  excluded: list[dict] | None = None) -> Path:
    """Write the results workbook. rows: one dict per cell and target (core.batch.result_rows)."""
    path = Path(path)
    rows = acquisition_notes(rows)
    wb = Workbook()
    used: set[str] = set()

    # Read me
    ws = wb.active
    ws.title = _sheet_title("Read me", used)
    ws.cell(row=1, column=1, value=f"Localization Tool {__version__}: results").font = Font(bold=True, size=14)
    n_cells = len({(r.get('experiment'), r['image'], r['cell']) for r in rows})
    ws.cell(row=2, column=1, value=f"Written {time.strftime('%Y-%m-%d %H:%M')}, {n_cells} cell(s), "
                                   f"{len(excluded or [])} excluded")
    r = 4
    for title, text in _read_me(info.get("params") or {}):
        c = ws.cell(row=r, column=1, value=title)
        if text is None:
            c.font = Font(bold=True, size=12, color=ACCENT)
        else:
            c.font = Font(bold=True)
            t = ws.cell(row=r, column=2, value=text)
            t.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    _widths(ws, {1: 28, 2: 110})

    # Summary
    _summary(wb.create_sheet(_sheet_title("Summary", used)), rows)

    # Prism sheets: checked cells only
    prism_rows, checked_only = _checked(rows)
    which = ("checked cells (accepted)" if checked_only else
             "ALL cells: none is checked yet, these values were computed automatically")
    targets = list(OrderedDict.fromkeys(r["target"] for r in prism_rows))
    conditions = sorted(OrderedDict.fromkeys((r.get("condition") or "no condition") for r in prism_rows))
    for target in targets:
        trows = _sorted([r for r in prism_rows if r["target"] == target])
        for key, short, explanation in KEY_MEASURES:
            if not any(isinstance(_num(r.get(key)), (int, float)) for r in trows):
                continue                     # e.g. no actin channel in any cell: no actin sheet
            ws = wb.create_sheet(_sheet_title(f"{target} {short}", used))
            ws.cell(row=1, column=1, value=f"{target}: {explanation}. {which}.").font = Font(bold=True)
            _header(ws, 2, conditions)
            fmt = PCT if key.endswith("_pct") else ENR
            id_col = len(conditions) + 2              # an empty column keeps the Prism block separate
            for c, cond in enumerate(conditions, start=1):
                sel = [r for r in trows if (r.get("condition") or "no condition") == cond]
                _header_cell(ws, 2, id_col + c - 1, f"ID {cond}", strong=False)
                for i, r in enumerate(sel, start=3):
                    ws.cell(row=i, column=c, value=_num(r.get(key))).number_format = fmt
                    ws.cell(row=i, column=id_col + c - 1,
                            value=_num(f"{r.get('experiment', '')} / {r.get('image', '')} / {r.get('cell', '')}"))
            _widths(ws, {c: 16 for c in range(1, len(conditions) + 1)})
            _widths(ws, {id_col + c: 40 for c in range(len(conditions))})
            ws.freeze_panes = "A3"

    # Replicates
    ws = wb.create_sheet(_sheet_title("Replicates", used))
    head = ["Target", "Experiment", "Condition", "Cells"] + [h for _, h, _ in KEY_MEASURES]
    _header(ws, 1, head)
    groups: "OrderedDict[tuple, list]" = OrderedDict()
    for row in _sorted(prism_rows):
        groups.setdefault((row["target"], row.get("experiment", ""), row.get("condition", "")), []).append(row)
    for i, ((target, exp, cond), grp) in enumerate(groups.items(), start=2):
        for c, v in enumerate((target, exp, cond, len(grp)), start=1):
            ws.cell(row=i, column=c, value=_num(v))
        for c, (key, _, _) in enumerate(KEY_MEASURES, start=5):
            vals = [v for v in (_num(g.get(key)) for g in grp) if isinstance(v, (int, float))]
            cell = ws.cell(row=i, column=c, value=sum(vals) / len(vals) if vals else None)
            cell.number_format = PCT if key.endswith("_pct") else ENR
    ws.cell(row=len(groups) + 3, column=1, value=f"Means of {which}.").font = Font(italic=True)
    _widths(ws, {c: 13 for c in range(1, len(head) + 1)})
    ws.freeze_panes = "E2"

    # Excluded
    if excluded:
        ws = wb.create_sheet(_sheet_title("Excluded", used))
        cols = ["image", "experiment", "condition", "cell", "reason"]
        _header(ws, 1, [c.capitalize() for c in cols])
        for i, row in enumerate(excluded, start=2):
            for c, k in enumerate(cols, start=1):
                ws.cell(row=i, column=c, value=_num(row.get(k)))
        _widths(ws, {1: 40, 2: 14, 3: 16, 4: 8, 5: 60})

    # All values
    ws = wb.create_sheet(_sheet_title("All values", used))
    cols: list[str] = []
    for row in rows:
        for k in row:
            if k not in cols and not k.startswith("_"):
                cols.append(k)
    _header(ws, 1, cols)
    for i, row in enumerate(_sorted(rows), start=2):
        for c, k in enumerate(cols, start=1):
            v = row.get(k)
            v = _num(STATUS_TEXT.get(v, v) if k == "status" else v)
            ws.cell(row=i, column=c, value=v if not isinstance(v, (list, dict)) else json.dumps(v))
    ws.freeze_panes = "F2"

    # Profiles
    ws = wb.create_sheet(_sheet_title("Profiles", used))
    pcols = ["image", "experiment", "condition", "cell", "target", "profile",
             "bin_start_um", "bin_end_um", "signal_fraction", "volume_fraction", "enrichment"]
    _header(ws, 1, pcols)
    for i, row in enumerate(profile_rows, start=2):
        for c, k in enumerate(pcols, start=1):
            ws.cell(row=i, column=c, value=_num(row.get(k)))

    # Acquisition (D63)
    acq_cols = sorted({k for row in rows for k in row if k.startswith(ACQ_PREFIXES)})
    if acq_cols:
        ws = wb.create_sheet(_sheet_title("Acquisition", used))
        modes = sorted({k for row in rows for k in row if k.startswith("airyscan_mode_")})
        head = ["experiment", "condition", "image"] + modes + acq_cols + ["differs within experiment"]
        _header(ws, 1, head)
        seen = set()
        i = 2
        for row in _sorted(rows):
            key = (row.get("experiment"), row.get("condition"), row.get("image"))
            if key in seen:
                continue
            seen.add(key)
            for c, k in enumerate(head[:-1], start=1):
                ws.cell(row=i, column=c, value=_num(row.get(k)))
            ws.cell(row=i, column=len(head), value=row.get("_acquisition_differs", ""))
            i += 1
        _widths(ws, {1: 14, 2: 16, 3: 40, len(head): 60})
        ws.freeze_panes = "D2"

    # Info
    ws = wb.create_sheet(_sheet_title("Info", used))
    r = 1
    for key, value in info.items():
        ws.cell(row=r, column=1, value=str(key)).font = Font(bold=True)
        text = value if isinstance(value, (str, int, float)) else json.dumps(value, indent=1, default=str)
        c = ws.cell(row=r, column=2, value=_num(text))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    _widths(ws, {1: 22, 2: 120})

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    try:
        wb.save(tmp)
        _replace(tmp, path)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
    return path


def _replace(tmp: Path, path: Path, tries: int = 5) -> None:
    """Windows virus scanners hold new files for a moment: try the replace a few times."""
    for k in range(tries):
        try:
            tmp.replace(path)
            return
        except PermissionError:
            if k == tries - 1:
                raise
            time.sleep(0.3)
