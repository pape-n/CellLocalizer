"""Methods text in English and German (D58).

Built from an analysis (numbers and settings as used) or from the default settings alone. It is a
complete methods section and at the same time the description of how the program computes: the
user guide points here. Parts the program cannot know (antibodies, microscope settings that are
not in the file) appear in [square brackets]. Every statement that rests on literature carries the
reference, the list at the end gives the DOI.
"""

from __future__ import annotations

import platform
from pathlib import Path
from collections import Counter
from dataclasses import dataclass, field

from .. import REPOSITORY, __version__
from .params import AnalysisParams

LANGUAGES = {"en": "English", "de": "Deutsch"}
AUTHOR = "Nils Pape"
CONTACT = "npape2@uni-muenster.de"

REFERENCES = {
    "li": "Li CH, Tam PKS (1998) An iterative algorithm for minimum cross entropy thresholding. Pattern Recognit "
          "Lett 19:771-780. doi:10.1016/S0167-8655(98)00057-9",
    "otsu": "Otsu N (1979) A threshold selection method from gray-level histograms. IEEE Trans Syst Man Cybern "
            "9:62-66. doi:10.1109/TSMC.1979.4310076",
    "triangle": "Zack GW, Rogers WE, Latt SA (1977) Automatic measurement of sister chromatid exchange frequency. "
                "J Histochem Cytochem 25:741-753. doi:10.1177/25.7.70454",
    "huff": "Huff J (2015) The Airyscan detector from ZEISS: confocal imaging with improved signal-to-noise ratio "
            "and super-resolution. Nat Methods 12:i-ii. doi:10.1038/nmeth.f.388",
    "diel": "Diel EE, Lichtman JW, Richardson DS (2020) Tutorial: avoiding and correcting sample-induced spherical "
            "aberration artifacts in 3D fluorescence microscopy. Nat Protoc 15:2773-2784. doi:10.1038/s41596-020-0360-2",
    "biermann": "Biermann H et al. (1999) Murine leukocytes with ring-shaped nuclei include granulocytes, "
                "monocytes, and their precursors. J Leukoc Biol 65:217-231. PMID 10088605",
    "carpenter": "Carpenter AE et al. (2006) CellProfiler: image analysis software for identifying and quantifying "
                 "cell phenotypes. Genome Biol 7:R100. doi:10.1186/gb-2006-7-10-r100",
    "ollion": "Ollion J et al. (2013) TANGO: a generic tool for high-throughput 3D image analysis for studying "
              "nuclear organization. Bioinformatics 29:1840-1841. doi:10.1093/bioinformatics/btt276",
    "costes": "Costes SV et al. (2004) Automatic and quantitative measurement of protein-protein colocalization in "
              "live cells. Biophys J 86:3993-4003. doi:10.1529/biophysj.103.038422",
    "manders": "Manders EMM, Verbeek FJ, Aten JA (1993) Measurement of co-localization of objects in dual-colour "
               "confocal images. J Microsc 169:375-382. doi:10.1111/j.1365-2818.1993.tb03313.x",
    "dunn": "Dunn KW, Kamocka MM, McDonald JH (2011) A practical guide to evaluating colocalization in biological "
            "microscopy. Am J Physiol Cell Physiol 300:C723-C742. doi:10.1152/ajpcell.00462.2010",
    "giese": "Giese W et al. (2025) Polarity-JaM: an image analysis toolbox for cell polarity, junction and "
             "morphology quantification. Nat Commun 16:1474. doi:10.1038/s41467-025-56643-x",
    "zimmermann": "Zimmermann T (2005) Spectral imaging and linear unmixing in light microscopy. Adv Biochem Eng "
                  "Biotechnol 95:245-265. doi:10.1007/b102216",
    "kozubek": "Kozubek M, Matula P (2000) An efficient algorithm for measurement and correction of chromatic "
               "aberrations in fluorescence microscopy. J Microsc 200:206-217. doi:10.1046/j.1365-2818.2000.00754.x",
    "matsuda": "Matsuda A et al. (2018) Accurate and fiducial-marker-free correction for three-dimensional "
               "chromatic shift in biological fluorescence microscopy. Sci Rep 8:7583. "
               "doi:10.1038/s41598-018-25922-7",
    "lord": "Lord SJ et al. (2020) SuperPlots: communicating reproducibility and variability in cell biology. "
            "J Cell Biol 219:e202001064. doi:10.1083/jcb.202001064",
    "schmied": "Schmied C et al. (2024) Community-developed checklists for publishing images and image analyses. "
               "Nat Methods 21:170-181. doi:10.1038/s41592-023-01987-9",
    "napari": "napari contributors (2019) napari: a multi-dimensional image viewer for Python. "
              "doi:10.5281/zenodo.3555620",
    "skimage": "van der Walt S et al. (2014) scikit-image: image processing in Python. PeerJ 2:e453. "
               "doi:10.7717/peerj.453",
    "scipy": "Virtanen P et al. (2020) SciPy 1.0: fundamental algorithms for scientific computing in Python. "
             "Nat Methods 17:261-272. doi:10.1038/s41592-019-0686-2",
}


@dataclass
class Summary:
    """What the analysis contains, for the numbers in the text. Empty: text from the settings alone."""
    images: int = 0
    experiments: int = 0
    conditions: list[str] = field(default_factory=list)
    cells: int = 0
    accepted: int = 0
    to_check: int = 0
    excluded: Counter = field(default_factory=Counter)          # reason -> number of cells or images
    targets: list[str] = field(default_factory=list)
    roles: dict[str, str] = field(default_factory=dict)          # label -> role
    voxel_um: tuple[float, float, float] | None = None          # (x, y, z), median over cells
    objective: str = ""
    immersion: str = ""
    airyscan: dict[str, tuple[float, float]] = field(default_factory=dict)   # label -> (min, max) strength
    airyscan_modes: set = field(default_factory=set)            # e.g. {"Auto"} or {"Manual"} from the metadata
    stale: int = 0                     # cells computed with other settings than the current ones
    cut_top: int = 0                   # cells with the note "stack cut off at the top"
    bottom_bright: int = 0             # cells whose plane at the coverslip is already bright (D67, column bottom_plane_ratio)
    pixel_range: tuple[float, float] | None = None              # smallest and largest pixel size in µm
    acquisition_differs: list[str] = field(default_factory=list)   # experiments with differing settings (D63)
    acquisition_read: bool = False     # laser and detector settings were found in the image metadata
    own_factors: int = 0
    blind: bool = False
    revealed: int = 0


def summary_from_analysis(af) -> Summary:
    """Counts and acquisition values from an AnalysisFile with stored results."""
    s = Summary(blind=af.blind)
    s.experiments = len({im.experiment for im in af.images if im.experiment}) or (1 if af.images else 0)
    s.conditions = sorted({im.condition for im in af.images if im.condition})
    vox, strengths, rows = [], {}, []
    used_images = set()
    for im in af.images:
        if im.revealed:
            s.revealed += 1
        roles = im.assignments or af.panel or []
        for a in roles:
            if a.role != "Ignore":
                s.roles.setdefault(a.label, a.role)
            if a.role == "Target" and a.label not in s.targets:
                s.targets.append(a.label)
        for c in im.cells:
            if c.status == "excluded":
                s.excluded[c.note or "no reason given"] += 1
                continue
            s.cells += 1
            used_images.add(str(im.path))
            from .batch import needs_update

            if needs_update(c, af.params):
                s.stale += 1
            if str(c.status).startswith("accepted"):
                s.accepted += 1
            else:
                s.to_check += 1
            if c.overrides:
                s.own_factors += 1
            row = ((c.results or {}).get("rows") or [{}])[0]
            if row:
                rows.append(row)
            notes = str(row.get("notes") or "")
            s.cut_top += "stack cut off at the top" in notes
            ratio = row.get("bottom_plane_ratio")
            limit = float(((c.results or {}).get("params") or {}).get("bottom_fraction", 0.6))
            s.bottom_bright += isinstance(ratio, (int, float)) and ratio == ratio and ratio >= limit
            if row.get("voxel_x_um"):
                # the stored z size includes the axial scale (D23); the text names the recorded z step
                scale = float(((c.results or {}).get("params") or {}).get("axial_scale", 1.0) or 1.0)
                vox.append((row["voxel_x_um"], row["voxel_y_um"], row["voxel_z_um"] / scale))
            s.objective = s.objective or str(row.get("objective") or "")
            s.immersion = s.immersion or str(row.get("immersion") or "")
            if any(k.startswith(("laser_transmission_pct_", "detector_voltage_")) for k in row):
                s.acquisition_read = True
            for k, v in row.items():
                if k.startswith("airyscan_strength_") and isinstance(v, (int, float)):
                    strengths.setdefault(k[len("airyscan_strength_"):], []).append(float(v))
                if k.startswith("airyscan_mode_") and v:
                    s.airyscan_modes.add(str(v))
    if vox:
        vox.sort()
        s.voxel_um = vox[len(vox) // 2]
        s.pixel_range = (min(v[0] for v in vox), max(v[0] for v in vox))
    s.airyscan = {k: (min(v), max(v)) for k, v in strengths.items()}
    s.images = len(used_images)                 # images that contribute cells, not every image in the list
    from .export_excel import acquisition_notes

    s.acquisition_differs = sorted({r.get("experiment") or "" for r in acquisition_notes(rows)
                                    if r.get("_acquisition_differs")})
    return s


class _Refs:
    def __init__(self):
        self.order: list[str] = []

    def __call__(self, *keys: str) -> str:
        nums = []
        for k in keys:
            if k not in self.order:
                self.order.append(k)
            nums.append(self.order.index(k) + 1)
        return "[" + ", ".join(str(n) for n in nums) + "]"


def _version(dist: str, module: str | None = None) -> str:
    """Version of an installed package without importing it (napari must not be imported in core)."""
    try:
        from importlib.metadata import version

        return version(dist)
    except Exception:                       # e.g. frozen build without metadata
        if module:
            try:
                import importlib

                return str(importlib.import_module(module).__version__)
            except Exception:
                pass
    return ""


def _n(v: float, lang: str, digits: int = 2) -> str:
    text = f"{v:.{digits}f}"
    if digits > 0:
        text = text.rstrip("0").rstrip(".")
    return text.replace(".", ",") if lang == "de" else text


def _join(items: list[str], lang: str) -> str:
    items = [i for i in items if i]
    if len(items) < 2:
        return "".join(items)
    return ", ".join(items[:-1]) + (" and " if lang == "en" else " und ") + items[-1]


METHOD_NAME = {"en": {"li": "Li's minimum cross-entropy method", "otsu": "Otsu's method",
                      "triangle": "the triangle method", "manual": "a fixed threshold",
                      "halfmax": "the half-maximum criterion"},
               "de": {"li": "dem Minimum-Kreuzentropie-Verfahren von Li", "otsu": "dem Verfahren von Otsu",
                      "triangle": "dem Dreiecksverfahren", "manual": "einer festen Schwelle",
                      "halfmax": "dem Halbmaximum-Kriterium"}}
METHOD_REF = {"li": "li", "otsu": "otsu", "triangle": "triangle"}


def methods_text(params: AnalysisParams | None = None, summary: Summary | None = None, lang: str = "en") -> str:
    """The full methods section with references. Without summary the counts appear as placeholders."""
    if lang not in LANGUAGES:
        raise ValueError(f"language must be one of {list(LANGUAGES)}")
    p = params or AnalysisParams()
    s = summary or Summary()
    ref = _Refs()
    en = lang == "en"
    n = lambda v, d=2: _n(float(v), lang, d)  # noqa: E731
    paras: list[tuple[str, str]] = []

    def fac_note(mp) -> str:
        if abs(mp.factor - 1.0) < 1e-9:
            return ""
        return (f" (threshold multiplied by {n(mp.factor)})" if en else f" (Schwelle multipliziert mit {n(mp.factor)})")

    def thr(mp, key: str) -> str:
        if mp.method == "manual":
            v = n(mp.manual_threshold or 0, 1)
            return f"a fixed threshold of {v}" if en else f"einer festen Schwelle von {v}"
        return ("a threshold by " if en else "einer Schwelle nach ") + meth(mp, key)

    def meth(mp, key: str) -> str:
        name = METHOD_NAME[lang][mp.method]
        r = f" {ref(METHOD_REF[mp.method])}" if mp.method in METHOD_REF else ""
        fac = "" if abs(mp.factor - 1.0) < 1e-9 else (f", multiplied by {n(mp.factor)}" if en else
                                                       f", multipliziert mit {n(mp.factor)}")
        return name + r + fac

    # ---- software
    py = platform.python_version()
    libs = {name: _version(dist, mod) for name, dist, mod in (("napari", "napari", None),
                                                               ("scikit-image", "scikit-image", "skimage"),
                                                               ("SciPy", "scipy", "scipy"), ("NumPy", "numpy", "numpy"))}

    def lib(name: str) -> str:
        return f"{name} {libs[name]}" if libs.get(name) else name
    if en:
        paras.append(("Software", f"Images were analysed with the Localization Tool, version {__version__} "
                      f"(developed by {AUTHOR}, Institute of Physiology II, University of Münster; open source, BSD-3-Clause, "
                      f"{REPOSITORY}), written in Python "
                      f"{py} with {lib('napari')} {ref('napari')}, {lib('scikit-image')} {ref('skimage')}, {lib('SciPy')} "
                      f"{ref('scipy')} and {lib('NumPy')}. "
                      "All settings and the program version are stored in the analysis file and in the results "
                      "workbook (sheet Info), so every value can be reproduced."
                      + (f" [{s.stale} cell(s) were computed with other settings than those described here: use "
                         "'Apply settings to all cells' before reporting.]" if s.stale else "")))
    else:
        paras.append(("Software", f"Die Bilder wurden mit dem Localization Tool, Version {__version__}, ausgewertet "
                      f"(entwickelt von {AUTHOR}, Institut für Physiologie II, Universität Münster; Open Source, BSD-3-Clause, "
                      f"{REPOSITORY}), geschrieben in "
                      f"Python {py} mit {lib('napari')} {ref('napari')}, {lib('scikit-image')} {ref('skimage')}, "
                      f"{lib('SciPy')} {ref('scipy')} und {lib('NumPy')}. Alle Einstellungen und die Programmversion sind in der Analysedatei und in "
                      "der Ergebnisdatei (Blatt Info) gespeichert, so dass jeder Wert reproduzierbar ist."
                      + (f" [{s.stale} Zelle(n) wurden mit anderen als den hier beschriebenen Einstellungen gerechnet: "
                         "vor dem Bericht 'Apply settings to all cells' ausführen.]" if s.stale else "")))

    # ---- acquisition
    obj = s.objective or ("[objective]" if en else "[Objektiv]")
    imm = s.immersion or ("[immersion]" if en else "[Immersion]")
    if s.voxel_um:
        vx, vy, vz = s.voxel_um
        rng_px = ""
        if s.pixel_range and s.pixel_range[1] - s.pixel_range[0] > 0.0015:
            lo_px, hi_px = (n(v * 1000, 0) for v in s.pixel_range)
            rng_px = f" ({lo_px} to {hi_px} nm)" if en else f" ({lo_px} bis {hi_px} nm)"
        vox = (f"a pixel size of {n(vx * 1000, 0)} nm{rng_px} and a z step of {n(vz, 3)} µm" if en else
               f"einer Pixelgröße von {n(vx * 1000, 0)} nm{rng_px} und einem z-Abstand von {n(vz, 3)} µm")
    else:
        vox = "a pixel size of [x] nm and a z step of [z] µm" if en else "einer Pixelgröße von [x] nm und einem z-Abstand von [z] µm"
    if s.airyscan:
        lo = min(v[0] for v in s.airyscan.values())
        hi = max(v[1] for v in s.airyscan.values())
        auto = bool(s.airyscan_modes) and all(m.lower().startswith("auto") for m in s.airyscan_modes)
        rng = n(lo, 1) if abs(hi - lo) < 0.05 else (f"{n(lo, 1)} to {n(hi, 1)}" if en else f"{n(lo, 1)} bis {n(hi, 1)}")
        if auto:
            strength = (f"with automatically estimated strength ({rng})" if en else
                        f"mit automatisch geschätzter Stärke ({rng})")
        elif not s.airyscan_modes:
            strength = (f"with strength {rng} [automatic or fixed]" if en else f"mit der Stärke {rng} [automatisch oder fest]")
        else:
            strength = f"with strength {rng}" if en else f"mit der Stärke {rng}"
    else:
        strength = "with strength [x]" if en else "mit der Stärke [x]"
    if s.acquisition_differs:
        exps = _join(s.acquisition_differs, lang)
        same = (f"Laser power, detector gain or Airyscan strength differed within experiment(s) {exps} (sheet "
                "Acquisition of the results workbook) [state why]." if en else
                f"Laserleistung, Detektorverstärkung oder Airyscan-Stärke unterschieden sich innerhalb von Experiment "
                f"{exps} (Blatt Acquisition der Ergebnisdatei) [Grund angeben].")
    elif s.cells and not s.acquisition_read:
        same = ("[Laser power and detector gain within an experiment: not in the image metadata, please state.]"
                if en else "[Laserleistung und Detektorverstärkung innerhalb eines Experiments: nicht in den "
                "Bildmetadaten, bitte angeben.]")
    elif s.cells:
        same = ("Laser power and detector gain were identical within each experiment according to the image "
                "metadata." if en else
                "Laserleistung und Detektorverstärkung waren laut Bildmetadaten innerhalb jedes Experiments gleich.")
    else:
        same = ("Laser power and detector gain were kept constant within an experiment; the program reports "
                "deviations read from the image metadata." if en else
                "Laserleistung und Detektorverstärkung waren innerhalb eines Experiments gleich; das Programm meldet "
                "Abweichungen aus den Metadaten.")
    if abs(p.axial_scale - 1.0) > 1e-9:
        vox += (f"; z distances were multiplied by {n(p.axial_scale, 3)} to correct the axial scaling caused by the "
                f"refractive index mismatch {ref('diel')}" if en else
                f"; z-Abstände wurden mit {n(p.axial_scale, 3)} multipliziert, um die axiale Verzerrung durch den "
                f"Brechungsindexunterschied zu korrigieren {ref('diel')}")
    complete = bool(s.cells) and not (s.cut_top or s.bottom_bright)
    if en:
        stacks = ("z-stacks covering the whole cell" if complete else
                  "z-stacks [from below the cell to above it]" if not s.cells else
                  "z-stacks (coverage of the cell: see quality control)")
        paras.append(("Image acquisition", f"Cells were imaged on a [Zeiss LSM 800 with Airyscan detector] "
                      f"{ref('huff')} with a {obj} objective ({imm}) as {stacks}, with {vox}. "
                      f"Raw data were processed with [3D] Airyscan processing in [ZEN version] {strength}. "
                      "[Laser lines, detection ranges, embedding medium.] Nuclei were stained with [DAPI], F-actin "
                      "with [phalloidin conjugate] and " + _targets(s, lang) + " with [primary antibody, clone, "
                      "dilution; secondary antibody]. " + same))
    else:
        stacks = ("z-Stapel über die ganze Zelle" if complete else
                  "z-Stapel [von unterhalb bis oberhalb der Zelle]" if not s.cells else
                  "z-Stapel (Abdeckung der Zelle: siehe Qualitätskontrolle)")
        paras.append(("Bildaufnahme", f"Die Zellen wurden an einem [Zeiss LSM 800 mit Airyscan-Detektor] "
                      f"{ref('huff')} mit einem {obj}-Objektiv ({imm}) als {stacks} "
                      f"aufgenommen, mit {vox}. Die Rohdaten wurden mit [3D-]Airyscan-Verarbeitung in [ZEN-Version] "
                      f"{strength} verarbeitet. [Laserlinien, Detektionsbereiche, Einbettmedium.] Kerne wurden mit "
                      "[DAPI], F-Aktin mit [Phalloidin-Konjugat] und " + _targets(s, lang) + " mit [Primärantikörper, "
                      "Klon, Verdünnung; Sekundärantikörper] gefärbt. " + same))

    # ---- preprocessing
    bg = {k: v for k, v in p.background.items() if v}
    if en:
        pre = ("Intensities were used as recorded; no background was subtracted, because Airyscan-processed data "
               "are clipped at zero." if not bg else
               "A constant background was subtracted per channel (" +
               _join([f"{k} {n(v, 1)}" for k, v in bg.items()], lang) + "), negative values were set to zero.")
        if p.crosstalk:
            pre += (" Spectral crosstalk was corrected linearly as I'_target = max(I_target - k I_source, 0) "
                    f"{ref('zimmermann')} with k estimated from single-stained controls (" +
                    _join([f"{c['into']} from {c['from']}: k = {n(c['k'], 3)}" for c in p.crosstalk], lang) + ").")
        if p.channel_offsets:
            pre += (" The chromatic offset between channels was measured with multicolour beads "
                    f"{ref('kozubek', 'matsuda')} and corrected by shifting channels by whole voxels (" +
                    _join([f"{o['channel']}: z {n(o['dz_um'], 3)}, y {n(o['dy_um'], 3)}, x {n(o['dx_um'], 3)} µm"
                           for o in p.channel_offsets], lang) + "); voxel values were not interpolated.")
        else:
            pre += (" No correction for chromatic offset between channels was applied [state the bead measurement "
                    "if done]; values of thin zones near surfaces depend on it.")
        paras.append(("Pre-processing", pre))
    else:
        pre = ("Die Intensitäten wurden wie aufgenommen verwendet; ein Hintergrund wurde nicht abgezogen, weil "
               "Airyscan-verarbeitete Daten bei null abgeschnitten sind." if not bg else
               "Pro Kanal wurde ein konstanter Hintergrund abgezogen (" +
               _join([f"{k} {n(v, 1)}" for k, v in bg.items()], lang) + "), negative Werte wurden auf null gesetzt.")
        if p.crosstalk:
            pre += (" Spektrales Übersprechen wurde linear korrigiert, I'_Ziel = max(I_Ziel - k I_Quelle, 0) "
                    f"{ref('zimmermann')}, mit k aus Einzelfärbungen (" +
                    _join([f"{c['into']} aus {c['from']}: k = {n(c['k'], 3)}" for c in p.crosstalk], lang) + ").")
        if p.channel_offsets:
            pre += (" Der chromatische Versatz zwischen den Kanälen wurde mit Mehrfarben-Beads gemessen "
                    f"{ref('kozubek', 'matsuda')} und durch Verschieben um ganze Voxel korrigiert (" +
                    _join([f"{o['channel']}: z {n(o['dz_um'], 3)}, y {n(o['dy_um'], 3)}, x {n(o['dx_um'], 3)} µm"
                           for o in p.channel_offsets], lang) + "); Voxelwerte wurden nicht interpoliert.")
        else:
            pre += (" Ein chromatischer Versatz zwischen den Kanälen wurde nicht korrigiert [Bead-Messung angeben, "
                    "falls erfolgt]; Werte dünner Zonen nahe Oberflächen hängen davon ab.")
        paras.append(("Vorverarbeitung", pre))

    # ---- segmentation
    c, nu, ac, ar, tg = p.cell, p.nucleus, p.actin, p.actin_rich, p.target
    halfmax_cell = c.method == "halfmax"
    second = c.method != "manual"
    if en:
        if halfmax_cell:
            edge = (f"first with Li's method {ref('li')} (σ = {n(c.sigma_um)} µm), computed on the image and again on "
                    "the box around the first cell plus 1 µm. This Li cell lies at the outer end of the blur. Each "
                    "channel's threshold was then set at half maximum, halfway between its background (median of the "
                    "voxels at least 1 µm outside the Li cell, within the box around it extended by 3 µm) and its level just inside the edge "
                    "(median of the voxels 0.3 to 0.8 µm inside the Li edge that lie above the channel's Li threshold; "
                    "a channel that hardly reaches the edge, such as DAPI, used its median above the Li threshold "
                    "inside the cell), where the edge of a blurred object lies" + fac_note(c) + ". These thresholds "
                    "were only allowed to trim the outer 0.5 µm of the Li cell (depth from its envelope, with cracks "
                    "narrower than 0.3 µm sealed and both stack ends counted as closed); deeper voxels remained cell "
                    "even where every stain was dim, e.g. cytoplasm between bright granules")
            fill = "holes filled in every plane and in 3D (both stack ends counted as closed)"
        else:
            edge = (f"with {meth(c, 'cell')} (σ = {n(c.sigma_um)} µm)"
                    + (", computed on the image and again on the box around the first cell plus 1 µm" if second else ""))
            fill = "holes filled in every plane and in 3D"
        cell = ("Segmentation was done in 3D on the voxel grid of the images, without resampling. Each mask was "
                "obtained from a Gaussian-smoothed channel (σ in µm, converted per axis) with one global threshold per "
                "channel. Cell: channels without staining above noise (99.9th percentile less than ten noise widths above "
                "the median; noise width from the differences of neighbouring voxels, for data clipped at zero from the "
                "median absolute deviation) were left out; every other channel was thresholded on its own " + edge
                + ". The channel masks, each without specks below 0.5 µm³, were combined, " + fill + ", regions "
                "excluded by hand (polygons through all planes) removed and the largest connected object kept "
                f"(minimum {n(c.min_volume_um3, 1)} µm³).")
        nuc = (f"Nucleus: the DAPI channel (σ = {n(nu.sigma_um)} µm) was thresholded inside the cell "
               + ("at half maximum between the median nuclear and the median cytoplasmic DAPI level (cytoplasm: cell "
                  "voxels at least 0.5 µm from the nucleus, or the 10th percentile outside the nucleus when the "
                  "nucleus left almost no such voxels), iterated from Otsu's threshold " + ref("otsu") +
                  " until it changed by less than 0.5 % (at most six steps)" + fac_note(nu)
                  if nu.method == "halfmax" else "with " + meth(nu, "nucleus"))
               + f". Holes up to {n(p.nucleus_max_hole_um2)} µm² per plane and cavities closed in 3D were filled "
               "(the stack end away from the coverslip counted as closed, the coverslip side as open, and cracks "
               "narrower than 0.4 µm were sealed for this test), whereas tunnels open to the cytoplasm were kept, so "
               f"that ring-shaped nuclei of mouse granulocytes {ref('biermann')} keep the cytoplasm in their centre. "
               + (f"Holes of any size in a plane whose mean smoothed DAPI signal reached {n(p.nucleus_dim_hole_fraction)} "
                  "times the threshold (dim euchromatin) were also counted as nucleus. "
                  if p.nucleus_dim_hole_fraction > 0 else "")
               + f"All nuclear parts above {n(nu.min_volume_um3, 1)} µm³ were kept. A DAPI channel without staining "
               "above noise gave no nucleus.")
        act = ("Actin: " + ("the shell within " + n(p.cortex_um) + " µm below the cell surface, outside the nucleus."
                            if p.actin_definition == "shell" else
                            f"phalloidin signal above {thr(ac, 'actin')}, computed over all cell voxels (σ = "
                            f"{n(ac.sigma_um)} µm)"
                            + (", or above the threshold of its own plane, set halfway between the median cytoplasmic "
                               "and the median actin signal of that plane and iterated like the nuclear threshold, so "
                               "that the dimmer cortex of upper planes was included also where the actin at the "
                               "coverslip is several times brighter" if p.actin_per_plane and ac.method != "manual"
                               else "")
                            + "; the actin compartment is its part outside the nucleus. The "
                            f"brightest part, above {thr(ar, 'actin_rich')} (σ = {n(ar.sigma_um)} µm) and within "
                            "the actin mask, was reported as actin-rich zone."))
        tmask = (f"For visual control, a target mask was drawn with {thr(tg, 'target')} (σ = {n(tg.sigma_um)} µm); "
                 "it was not formed for channels without staining above noise, and the main values do not depend on it.")
        fac = (f" For {s.own_factors} cell(s) a threshold factor was adjusted after visual inspection; the factors are "
               "listed per cell in the workbook." if s.own_factors else
               " All cells were analysed with the same settings.")
        paras.append(("Segmentation", " ".join([cell, nuc, act, tmask]) + fac))
    else:
        if halfmax_cell:
            edge = (f"zuerst mit dem Verfahren nach Li {ref('li')} (σ = {n(c.sigma_um)} µm), berechnet auf dem Bild und "
                    "erneut auf dem Kasten um die erste Zelle plus 1 µm. Diese Li-Zelle liegt am äußeren Ende der "
                    "Unschärfe. Danach wurde die Schwelle jedes Kanals auf das Halbmaximum gesetzt, in die Mitte "
                    "zwischen seinem Hintergrund (Median der Voxel mindestens 1 µm außerhalb der Li-Zelle, im um 3 µm "
                    "erweiterten Kasten um sie) und seinem Niveau direkt innerhalb des Randes (Median der Voxel 0,3 bis 0,8 µm innerhalb "
                    "des Li-Randes, die über der Li-Schwelle des Kanals liegen; ein Kanal, der den Rand kaum erreicht, "
                    "etwa DAPI, nahm seinen Median über der Li-Schwelle in der Zelle), wo der Rand eines unscharf "
                    "abgebildeten Objekts liegt" + fac_note(c) + ". Diese Schwellen durften nur die äußeren 0,5 µm der "
                    "Li-Zelle beschneiden (Tiefe ab ihrer Hülle, Spalten unter 0,3 µm geschlossen, beide Stapelenden "
                    "als geschlossen gezählt); tiefere Voxel blieben Zelle, auch wo alle Färbungen schwach waren, etwa "
                    "Zytoplasma zwischen hellen Granula")
            fill = "Löcher in jeder Ebene und in 3D gefüllt (beide Stapelenden als geschlossen gezählt)"
        else:
            edge = (f"mit {meth(c, 'cell')} (σ = {n(c.sigma_um)} µm)"
                    + (", berechnet auf dem Bild und erneut auf dem Kasten um die erste Zelle plus 1 µm" if second else ""))
            fill = "Löcher in jeder Ebene und in 3D gefüllt"
        cell = ("Segmentiert wurde in 3D im Voxelraster der Bilder, ohne Neuabtastung. Jede Maske entstand aus einem "
                "gaußgeglätteten Kanal (σ in µm, pro Achse umgerechnet) mit einer globalen Schwelle pro Kanal. Zelle: "
                "Kanäle ohne Färbung über dem Rauschen (99,9. Perzentil weniger als zehn Rauschbreiten über dem "
                "Median; Rauschbreite aus den Differenzen benachbarter Voxel, bei auf null abgeschnittenen Daten aus "
                "der medianen absoluten Abweichung) wurden ausgelassen; jeder andere Kanal wurde einzeln geschwellt, " + edge + ". Die "
                "Kanalmasken, jeweils ohne Flecken unter 0,5 µm³, wurden vereinigt, " + fill + ", von Hand "
                "ausgeschlossene Bereiche (Polygone durch alle Ebenen) entfernt und das größte zusammenhängende "
                f"Objekt behalten (mindestens {n(c.min_volume_um3, 1)} µm³).")
        nuc = (f"Kern: der DAPI-Kanal (σ = {n(nu.sigma_um)} µm) wurde innerhalb der Zelle "
               + ("beim Halbmaximum zwischen medianem Kern- und Zytoplasmaniveau geschwellt (Zytoplasma: Zellvoxel "
                  "mindestens 0,5 µm vom Kern, oder das 10. Perzentil außerhalb des Kerns, wenn der Kern kaum solche "
                  "Voxel übrig ließ), iteriert ab der Otsu-Schwelle " + ref("otsu") +
                  ", bis sie sich um weniger als 0,5 % änderte (höchstens sechs Schritte)" + fac_note(nu)
                  if nu.method == "halfmax" else "mit " + meth(nu, "nucleus") + " geschwellt")
               + f". Löcher bis {n(p.nucleus_max_hole_um2)} µm² pro Ebene und in 3D geschlossene Hohlräume wurden "
               "gefüllt (das Stapelende weg vom Deckglas galt als geschlossen, die Deckglasseite als offen, Spalten "
               "unter 0,4 µm wurden für diese Prüfung geschlossen), Tunnel zum Zytoplasma blieben offen, so dass "
               f"ringförmige Kerne muriner Granulozyten {ref('biermann')} das Zytoplasma in ihrer Mitte behalten. "
               + (f"Löcher beliebiger Größe in einer Ebene, deren mittleres geglättetes DAPI-Signal mindestens das "
                  f"{n(p.nucleus_dim_hole_fraction)}-fache der Schwelle erreichte (schwaches Euchromatin), zählten "
                  "ebenfalls zum Kern. " if p.nucleus_dim_hole_fraction > 0 else "")
               + f"Alle Kernteile über {n(nu.min_volume_um3, 1)} µm³ wurden behalten. Ein DAPI-Kanal ohne Färbung "
               "über dem Rauschen ergab keinen Kern.")
        act = ("Aktin: " + ("die Schale bis " + n(p.cortex_um) + " µm unter der Zelloberfläche, außerhalb des Kerns."
                            if p.actin_definition == "shell" else
                            f"Phalloidinsignal über {thr(ac, 'actin')}, berechnet über alle Zellvoxel (σ = "
                            f"{n(ac.sigma_um)} µm)"
                            + (", oder über der Schwelle der eigenen Ebene, gesetzt auf halbem Weg zwischen medianem "
                               "Zytoplasma- und medianem Aktinsignal dieser Ebene und iteriert wie die Kernschwelle, so "
                               "dass der schwächere Kortex oberer Ebenen auch dort erfasst wurde, wo das Aktin am "
                               "Deckglas um ein Mehrfaches heller ist" if p.actin_per_plane and ac.method != "manual"
                               else "")
                            + "; das Aktinkompartiment ist sein Teil außerhalb des Kerns. Der "
                            f"hellste Teil, über {thr(ar, 'actin_rich')} (σ = {n(ar.sigma_um)} µm) und innerhalb "
                            "der Aktinmaske, wurde als aktinreiche Zone ausgegeben."))
        tmask = (f"Zur Sichtkontrolle wurde eine Target-Maske mit {thr(tg, 'target')} (σ = {n(tg.sigma_um)} µm) "
                 "gebildet, nicht bei Kanälen ohne Färbung über dem Rauschen; die Hauptwerte hängen nicht von ihr ab.")
        fac = (f" Bei {s.own_factors} Zelle(n) wurde ein Schwellenfaktor nach Sichtprüfung angepasst; die Faktoren "
               "stehen pro Zelle in der Ergebnisdatei." if s.own_factors else
               " Alle Zellen wurden mit denselben Einstellungen ausgewertet.")
        paras.append(("Segmentierung", " ".join([cell, nuc, act, tmask]) + fac))

    # ---- measures
    env, cor, bas = n(p.envelope_um), n(p.cortex_um), n(p.basal_um)
    if en:
        meas = ("Nucleus and cytoplasm (all cell voxels outside the nucleus) partition the cell; actin is the part "
                "of the cytoplasm with actin signal. For each target, the signal share of a region is the sum of its "
                "voxel intensities divided by the sum over the cell (f), the volume share is its number of voxels "
                "divided by that of the cell (φ), as in common object-based measurements "
                f"{ref('carpenter')}, and the enrichment is E = f / φ (1: as dense as the cell average). In addition, "
                f"zones were measured with exact Euclidean distance maps on the anisotropic voxel grid "
                f"{ref('ollion')}: nuclear envelope zone (within {env} µm of the nuclear border, about the axial "
                f"resolution), nucleus interior (more than {env} µm inside), cortex shell ({cor} µm below the cell "
                f"surface; the faces of the stack, including a cut top, count as surface) and basal layer ({bas} µm "
                "above the lowest cell plane, on the coverslip side taken from the image metadata). Zones overlap the "
                "compartments and each other.")
        if p.sensitivity:
            meas += (" To show how much each value depends on the exact edges, the signal shares and enrichments of "
                     "nucleus, cytoplasm and actin were recomputed with the cell threshold, and separately the nucleus "
                     "threshold, multiplied by 0.8 and 1.2 (the other edges and the actin threshold unchanged) and "
                     "reported as range; zones were not recomputed.")
        if p.polarity:
            meas += (" Polarity: an axis in the coverslip plane from the cell centroid to the actin-weighted centroid "
                     f"{ref('giese')}; polarity index = axis length / radius of the cell footprint; the plane "
                     "through the centroid perpendicular to the axis defined the front half.")
        if p.colocalization:
            meas += (" For comparison with the literature, Pearson's correlation coefficient and Manders' "
                     f"coefficients {ref('manders')} with the automatic threshold of Costes et al. {ref('costes')} "
                     f"were computed within the cell {ref('dunn')}; the Costes threshold exists only for positively "
                     "correlated channels, otherwise the Manders values were left empty.")
        paras.append(("Quantification", meas))
    else:
        meas = ("Kern und Zytoplasma (alle Zellvoxel außerhalb des Kerns) teilen die Zelle; Aktin ist der Teil des "
                "Zytoplasmas mit Aktinsignal. Für jedes Target ist der Signalanteil einer Region die Summe ihrer "
                "Voxelintensitäten geteilt durch die Summe über die Zelle (f), der Volumenanteil ihre Voxelzahl "
                "geteilt durch die der Zelle (φ), wie bei üblichen objektbasierten Messungen "
                f"{ref('carpenter')}, und die Anreicherung E = f / φ (1: so dicht wie im Mittel der Zelle). Zusätzlich "
                f"wurden Zonen mit exakten euklidischen Distanzkarten im anisotropen Voxelraster gemessen "
                f"{ref('ollion')}: Kernhüllenzone (bis {env} µm um den Kernrand, etwa die axiale Auflösung), "
                f"Kerninneres (mehr als {env} µm innen), Kortexschale ({cor} µm unter der Zelloberfläche; die Flächen "
                f"des Stapels, auch ein abgeschnittenes oberes Ende, zählen als Oberfläche) und basale Schicht ({bas} µm "
                "über der untersten Zellebene, auf der Deckglasseite laut Metadaten). Zonen überlappen die "
                "Kompartimente und einander.")
        if p.sensitivity:
            meas += (" Um zu zeigen, wie stark jeder Wert von den genauen Rändern abhängt, wurden Signalanteile und "
                     "Anreicherungen von Kern, Zytoplasma und Aktin mit dem 0,8- und 1,2-Fachen der Zellschwelle und "
                     "getrennt davon der Kernschwelle erneut berechnet (die übrigen Ränder und die Aktinschwelle "
                     "unverändert) und als Spanne angegeben; Zonen wurden nicht neu berechnet.")
        if p.polarity:
            meas += (" Polarität: eine Achse in der Deckglasebene vom Zellschwerpunkt zum aktingewichteten "
                     f"Schwerpunkt {ref('giese')}; Polaritätsindex = Achsenlänge / Radius der Zellgrundfläche; die "
                     "Ebene durch den Schwerpunkt senkrecht zur Achse definiert die Fronthälfte.")
        if p.colocalization:
            meas += (" Zum Vergleich mit der Literatur wurden der Korrelationskoeffizient nach Pearson und die "
                     f"Manders-Koeffizienten {ref('manders')} mit der automatischen Schwelle nach Costes et al. "
                     f"{ref('costes')} innerhalb der Zelle berechnet {ref('dunn')}; die Costes-Schwelle existiert nur "
                     "für positiv korrelierte Kanäle, sonst blieben die Manders-Werte leer.")
        paras.append(("Quantifizierung", meas))

    # ---- quality control
    if s.cells or s.excluded:
        exc = sum(s.excluded.values())
        reasons = _join([f"{r} ({k})" for r, k in s.excluded.most_common()], lang)
        if en:
            qc = (f"{s.cells} cell(s) from {s.images} image(s) and {s.experiments} experiment(s) were analysed; "
                  f"{s.accepted} were accepted after inspection of the masks in every plane and in 3D"
                  + (f", {s.to_check} were not yet inspected" if s.to_check else "")
                  + (f". {exc} cell(s) or image(s) were excluded: {reasons}." if exc else ". No cell was excluded."))
            if s.blind:
                qc += (" Cells were inspected blinded to file name and condition"
                       + (f" (the name was shown for {s.revealed} image(s))." if s.revealed else "."))
        else:
            qc = (f"{s.cells} Zelle(n) aus {s.images} Bild(ern) und {s.experiments} Experiment(en) wurden "
                  f"ausgewertet; {s.accepted} wurden nach Prüfung der Masken in allen Ebenen und in 3D angenommen"
                  + (f", {s.to_check} noch nicht geprüft" if s.to_check else "")
                  + (f". {exc} Zelle(n) oder Bild(er) wurden ausgeschlossen: {reasons}." if exc else
                     ". Keine Zelle wurde ausgeschlossen."))
            if s.blind:
                qc += (" Die Zellen wurden ohne Kenntnis von Dateiname und Bedingung geprüft"
                       + (f" (bei {s.revealed} Bild(ern) wurde der Name angezeigt)." if s.revealed else "."))
    else:
        qc = ("[n] cells from [n] images and [n] experiments were analysed; [n] were accepted after inspection of the "
              "masks, [n] were excluded ([reasons])." if en else
              "[n] Zellen aus [n] Bildern und [n] Experimenten wurden ausgewertet; [n] wurden nach Prüfung der Masken "
              "angenommen, [n] ausgeschlossen ([Gründe]).")
    if s.cells and not s.to_check:
        qc += (" The notes of the program (e.g. stack cut off at the top, possibly two cells) were reviewed for every "
               "accepted cell." if en else " Die Hinweise des Programms (z. B. Stapel oben abgeschnitten, "
               "möglicherweise zwei Zellen) wurden für jede angenommene Zelle geprüft.")
    else:
        qc += (" [State how the notes of the program (e.g. stack cut off at the top, possibly two cells) were "
               "handled.]" if en else " [Angeben, wie mit den Hinweisen des Programms (z. B. Stapel oben "
               "abgeschnitten, möglicherweise zwei Zellen) umgegangen wurde.]")
    if s.cut_top or s.bottom_bright:
        if en:
            parts = []
            if s.cut_top:
                parts.append(f"in {s.cut_top} of {s.cells} cell(s) the stack ended below the top of the cell, so the "
                             "cut face counts as cell surface")
            if s.bottom_bright:
                parts.append(f"in {s.bottom_bright} of {s.cells} cell(s) the plane at the coverslip was already at least "
                             f"{n(100 * p.bottom_fraction, 0)} % as bright as the brightest plane, so the stack probably "
                             "started inside the cell and the basal layer is the lowest recorded layer rather than the "
                             "layer on the coverslip")
            qc += " Coverage of the cells: " + "; ".join(parts) + "."
        else:
            parts = []
            if s.cut_top:
                parts.append(f"bei {s.cut_top} von {s.cells} Zelle(n) endete der Stapel unterhalb der Zelloberseite, die "
                             "Schnittfläche zählt dann als Zelloberfläche")
            if s.bottom_bright:
                parts.append(f"bei {s.bottom_bright} von {s.cells} Zelle(n) war die Ebene am Deckglas schon mindestens "
                             f"{n(100 * p.bottom_fraction, 0)} % so hell wie die hellste Ebene, der Stapel begann also "
                             "wahrscheinlich in der Zelle, und die basale Schicht ist die unterste aufgenommene Schicht, "
                             "nicht die Schicht auf dem Deckglas")
            qc += " Abdeckung der Zellen: " + "; ".join(parts) + "."
    paras.append(("Quality control" if en else "Qualitätskontrolle", qc))

    # ---- statistics
    if en:
        paras.append(("Statistics", "Values were computed per cell. Cells from one coverslip are not independent; "
                      "statistical comparisons used the mean per independent experiment [test, software, n], and "
                      f"single cells are shown together with the experiment means {ref('lord')}."))
        paras.append(("Data availability", "[Raw images, the analysis file (.lta.json) and the results workbook are "
                      "available from ...] The program is available from the author on request "
                      f"({AUTHOR}, {CONTACT})."))
    else:
        paras.append(("Statistik", "Die Werte wurden pro Zelle berechnet. Zellen eines Deckglases sind nicht "
                      "unabhängig; für statistische Vergleiche wurde der Mittelwert pro unabhängigem Experiment "
                      f"verwendet [Test, Software, n], und Einzelzellen werden zusammen mit den Experimentmitteln "
                      f"gezeigt {ref('lord')}."))
        paras.append(("Datenverfügbarkeit", "[Rohbilder, Analysedatei (.lta.json) und Ergebnisdatei sind erhältlich "
                      f"bei ...] Das Programm ist auf Anfrage beim Autor erhältlich ({AUTHOR}, {CONTACT})."))

    checklist = _checklist(s, p, lang, ref)
    lines = []
    for title, text in paras:
        lines += [title, text, ""]
    lines += checklist
    lines += ["References" if en else "Literatur"]
    lines += [f"{i}. {REFERENCES[k]}" for i, k in enumerate(ref.order, start=1)]
    return "\n".join(lines).strip() + "\n"


def _targets(s: Summary, lang: str) -> str:
    if s.targets:
        return (("the target protein " if len(s.targets) == 1 else "the target proteins ") if lang == "en" else
                ("das Zielprotein " if len(s.targets) == 1 else "die Zielproteine ")) + _join(s.targets, lang)
    return "the target protein [name]" if lang == "en" else "das Zielprotein [Name]"


def _checklist(s: Summary, p: AnalysisParams, lang: str, ref) -> list[str]:
    """What the community checklists ask for (Schmied et al. 2024), and whether this text covers it."""
    en = lang == "en"
    have, add = ("covered", "add") if en else ("enthalten", "ergänzen")
    items = [
        ("Software name, version and availability" if en else "Software, Version und Verfügbarkeit", have),
        ("Microscope, objective, detector and processing" if en else "Mikroskop, Objektiv, Detektor, Verarbeitung",
         ("partly, fill in the [brackets]" if en else "teilweise, [Klammern] ausfüllen") if s.objective else add),
        ("Labelling: antibodies, dyes, dilutions" if en else "Markierung: Antikörper, Farbstoffe, Verdünnungen", add),
        ("Every processing step with its parameters" if en else "Jeder Verarbeitungsschritt mit Parametern", have),
        ("Segmentation method and how it was checked" if en else "Segmentierung und ihre Prüfung", have),
        ("Definition of every reported measure" if en else "Definition jeder Messgröße", have),
        ("Inclusion and exclusion of cells" if en else "Ein- und Ausschluss von Zellen", have if s.cells else add),
        ("Example images with the masks (control images)" if en else "Beispielbilder mit Masken (Kontrollbilder)",
         add),
        ("Raw data and analysis files available" if en else "Rohdaten und Analysedateien verfügbar", add),
        ("Statistics: unit of replication, test" if en else "Statistik: Replikationseinheit, Test", add),
    ]
    head = ("Reporting checklist " + ref("schmied") if en else "Checkliste für die Veröffentlichung " + ref("schmied"))
    return [head] + [f"- {name}: {state}" for name, state in items] + [""]


def save_docx(text: str, path: Path) -> Path:
    """Methods text as Word file: section titles bold, references as a list (python-docx)."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    blocks = text.strip().split("\n\n")
    for block in blocks:
        lines = block.split("\n")
        if len(lines) >= 2 and not lines[0].startswith(("-", "1.")):
            doc.add_paragraph().add_run(lines[0]).bold = True
            for line in lines[1:]:
                doc.add_paragraph(line)
        else:
            for line in lines:
                doc.add_paragraph(line)
    doc.save(str(path))
    return path
