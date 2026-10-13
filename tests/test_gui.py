"""GUI test of the guided workflow (D17, D19, D21). Needs OpenGL: run with xvfb-run on Linux."""

import csv
import os
import sys
import time

from qtpy.QtWidgets import QFormLayout

import numpy as np
import pytest

pytestmark = pytest.mark.gui

if sys.platform.startswith("linux") and not os.environ.get("DISPLAY"):
    pytest.skip("no display (use xvfb-run)", allow_module_level=True)


def _pump(seconds=0.3):
    from qtpy.QtWidgets import QApplication

    t = time.time()
    while time.time() - t < seconds:
        QApplication.processEvents()
        time.sleep(0.01)


@pytest.fixture
def window(tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path))
    monkeypatch.setenv("XDG_CONFIG_HOME", str(tmp_path))
    import napari

    from localization_tool.gui import app as A
    from localization_tool.gui import theme

    viewer = napari.Viewer(show=True)
    viewer.theme = theme.register_napari_theme()
    w = A.MainWindow(viewer)
    w.setup()
    yield w
    viewer.close()


def test_reduced_interface(window):
    qv = window.viewer.window._qt_viewer
    assert not qv.dockLayerList.isVisible()
    assert not qv.dockLayerControls.isVisible()
    titles = [a.text() for a in window.viewer.window._qt_window.menuBar().actions() if a.isVisible()]
    assert titles == ["File", "Settings", "Help"]
    # step 1: nothing loaded, next is disabled
    assert window.panel.step == 0 and not window.panel.next_btn.isEnabled()


def _wait(w, seconds=30):
    t = time.time()
    while time.time() - t < seconds:
        _pump(0.05)
        if not w.state.computing:
            return
    raise AssertionError("calculation did not finish")


def test_guided_workflow_on_example(window, tmp_path):
    w = window
    p = w.panel
    w.load_example()
    assert p.step == 0                                   # opening stays in step 1, the user clicks Next (D62)
    assert p.next_btn.isEnabled()                        # example roles are preset
    layers_before = list(w.image_layers())
    for layer in layers_before:
        assert layer.interpolation2d == "linear"         # smooth display (D19)

    # changing a role updates the layers in place (0.2.0a1 took seconds here)
    a = [type(x)(**vars(x)) for x in w.state.assignments]
    a[1].role = "Ignore"
    t = time.perf_counter()
    w.apply_assignments(a)
    assert time.perf_counter() - t < 1.0
    assert w.image_layers() == layers_before and not layers_before[1].visible
    a[1].role = "Target"
    w.apply_assignments(a)

    # step 3: the cell outline is computed and shown
    p.go(2)
    _wait(w)
    assert w.state.cell is not None and w.state.cell.any()
    assert p.next_btn.isEnabled()
    cell_layer = w._layer(w.CELL_LAYER)
    assert cell_layer.visible and cell_layer.data.any()

    # remove a corner region: the cell is recomputed without it
    w.start_exclusion()
    w._layer(w.EXCLUDE_LAYER).add_polygons(np.array([[0, 0], [0, 60], [60, 0]], float))
    w.apply_exclusions()
    _wait(w)
    assert len(w.state.exclusions) == 1 and not w.state.cell[:, 0, 0].any()

    # edge factor: a lower factor gives a larger cell
    n_before = w.state.cell.sum()
    p.cell_factor.setValue(0.6)
    w.compute_cell(sync=True)
    assert w.state.cell.sum() > n_before
    p.cell_factor.setValue(1.0)
    w.compute_cell(sync=True)

    # in 3D the cell edge has no contour (contours exist only in 2D), back in 2D it has
    w.set_ndisplay(3)
    assert cell_layer.contour == 0
    w.set_ndisplay(2)
    assert cell_layer.contour == 1

    # step 4 computes in a worker with that cell, no image layers are added
    n_images = len(w.image_layers())
    p.go(3)
    _wait(w)
    r = w.state.result
    assert r is not None and len(w.image_layers()) == n_images
    assert np.array_equal(r.masks["cell"], w.state.cell)
    lab = w._layer(w.LABEL_LAYER)
    assert lab.visible and (lab.data == 4).any()
    assert set(w._target_layers) == {"TTN", "VIM"}
    p.front_box.setChecked(True)
    assert w._layer(w.FRONT_LAYER).visible
    # cytoplasm includes the actin (D39): with actin hidden, actin voxels get the cytoplasm colour
    p.show_box["actin"].setChecked(False)
    cmap = lab.colormap.color_dict
    assert tuple(cmap[2]) == tuple(cmap[3]) and cmap[3][3] > 0
    p.show_box["actin"].setChecked(True)
    # image channels can be hidden to see the masks alone
    boxes = [p.comp_channels.itemAt(k).widget() for k in range(p.comp_channels.count())]
    boxes = [b for b in boxes if b is not None and hasattr(b, "setChecked")]
    boxes[0].setChecked(False)
    assert not w.image_layers()[0].visible
    boxes[0].setChecked(True)
    # 3D: smoothed surfaces instead of voxel masks, computed in the background
    w.set_ndisplay(3)
    t = time.time()
    while w._surface_key is None and time.time() - t < 60:
        _pump(0.1)
    assert w._surfaces["nucleus"].visible and not lab.visible
    w.set_ndisplay(2)
    assert lab.visible and not w._surfaces["nucleus"].visible

    # show only the nucleus
    p.only_menu.setCurrentIndex(1)
    cmap = lab.colormap.color_dict
    assert cmap[2][3] == 0 and cmap[1][3] > 0
    p.only_menu.setCurrentIndex(5)
    assert w._target_layers["TTN"].visible
    p.only_menu.setCurrentIndex(0)

    for v in (0.9, 1.0):
        p.factor["actin"].setValue(v)
        w.compute(sync=True)
    assert len(w.image_layers()) == n_images

    p.go(4)
    assert p.results.rowCount() >= 20                    # two targets with zones and target mask
    w.accept_cell()
    assert sum(len(im.cells) for im in w.session.images) == 1 and not p.accept_btn.isEnabled()
    assert p.cell_list.count() == 1

    # save: one folder with the workbook, control images, masks and the analysis file
    out = tmp_path / "analysis"
    path = w.save_results(folder=out, notify=False)
    assert path.name == "analysis.lta.json"
    assert (out / "Results.xlsx").exists()
    assert len(list((out / "Control images").glob("*.png"))) == 1
    assert not (out / "masks").exists()                 # no painting any more, masks come from factors
    from openpyxl import load_workbook
    wb = load_workbook(out / "Results.xlsx")
    assert wb.sheetnames[0] == "Read me" and "TTN Enr nucleus" in wb.sheetnames
    summary = list(wb["Summary"].values)
    assert len(summary) == 4 and summary[2][5] == "accepted"      # two header rows

    # open the analysis again: the cell is listed without recalculation, and can be edited
    w.new_analysis()
    assert p.cell_list.count() == 0
    w.open_analysis(path)
    assert p.cell_list.count() == 1 and w.session_dir == out
    w.edit_cell(0)
    _wait(w)
    assert w.state.editing == "cell1" and p.step == 2 and w.state.cell is not None

    # exclude the cell with a reason: listed, not in the numbers, sheet Excluded
    p.go(3)
    _wait(w)
    p.go(4)
    w.exclude_cell("Out of focus")
    assert w.session.images[0].cells[0].status == "excluded"
    assert "Out of focus" in p.cell_list.item(0).text()
    w.save_results(folder=out, notify=False)
    wb = load_workbook(out / "Results.xlsx")
    assert "Excluded" in wb.sheetnames and len(list(wb["Summary"].values)) == 2


def test_performance_test_still_runs(window):
    from localization_tool.gui.perf import run_performance_test

    window.load_example()
    report = run_performance_test(window)
    assert "3D rotation" in report


def _golden_file(path, shift):
    from localization_tool.core.reader import write_imagej_tiff
    from localization_tool.core.synthetic import make_golden_cell

    stack, _ = make_golden_cell(voxel_um=(0.13, 0.08, 0.08), nucleus_shift_x=shift, margin_um=1.5)
    write_imagej_tiff(stack, path)
    return path


def _wait_batch(w, seconds=120):
    t = time.time()
    while time.time() - t < seconds:
        _pump(0.05)
        if w._batch is None:
            return
    raise AssertionError("batch did not finish")


def test_settings_and_all_cells(window, tmp_path, monkeypatch):
    from localization_tool.core.channels import ChannelAssignment
    from localization_tool.core.params import AnalysisParams
    from localization_tool.gui.settings import SettingsDialog

    w = window
    monkeypatch.setattr(w, "message", lambda *a, **k: None)
    monkeypatch.setattr(w, "ask", lambda *a, **k: False)
    folder = tmp_path / "20260301 - WT treated"
    for sub in ("Cell1", "Cell2"):
        (folder / sub).mkdir(parents=True)
    _golden_file(folder / "Cell1" / "a.tif", -1.0)
    _golden_file(folder / "Cell2" / "b.tif", -0.5)
    (folder / "Cell1" / "notes.txt").write_text("not an image")
    w.add_folder(folder, sync=True)                                  # a whole day at once
    assert [im.path.name for im in w.session.images] == ["a.tif", "b.tif"]
    assert w.session.images[0].experiment == "2026-03-01"           # suggestions from the folder (D37)
    assert w.session.images[0].condition == "WT treated"
    roles = [ChannelAssignment("Target", "T", "#FF2BD6"), ChannelAssignment("Actin", "Actin", "#2BD46B"),
             ChannelAssignment("Nucleus", "Nucleus", "#3D7DFF")]
    for im in w.session.images:
        im.assignments = [ChannelAssignment(**vars(a)) for a in roles]

    # analyse all images: one automatic cell each, with results and a control image waiting
    w.run_all("new")
    _wait_batch(w)
    cells = [c for im in w.session.images for c in im.cells]
    assert len(cells) == 2 and all(c.status == "automatic" and c.results["rows"] for c in cells)
    assert len(w.pending) == 2
    assert "to check" in w.panel.image_list.item(0).text()

    # new settings: cells computed with other settings are found and computed again
    d = AnalysisParams().to_dict()
    d["nucleus"]["factor"] = 0.9
    w._set_params(AnalysisParams.from_dict(d), ask=False)
    w.run_all("all")
    _wait_batch(w)
    cells = [c for im in w.session.images for c in im.cells]
    assert all(c.results["params"]["nucleus"]["factor"] == 0.9 for c in cells)
    w.run_all("all")                                                 # nothing changed: nothing to do
    _wait_batch(w)

    # 'Next image' opens the first cell to check
    w.next_image()
    t = time.time()
    while w.state.result is None and time.time() - t < 60:
        _pump(0.1)
    assert w.state.editing == "cell1" and w.panel.factors()["nucleus"] == pytest.approx(0.9)

    # the settings dialog shows and returns the same values
    dlg = SettingsDialog(None, w.params, ["T", "Actin", "Nucleus"], w.session.vocabulary)
    params, vocab = dlg.values()
    assert params.to_dict() == w.params.to_dict() and vocab == w.session.vocabulary
    dlg._add_xt("T", "Actin", 0.05)
    params, _ = dlg.values()
    assert params.crosstalk == [{"into": "T", "from": "Actin", "k": 0.05}]


def test_older_save_never_overwrites_a_newer_one(tmp_path):
    from localization_tool.gui import app as A

    def snap(seq, text):
        return {"folder": tmp_path, "seq": seq, "pending": {}, "json": text, "rows": [], "prof": [],
                "excluded": [], "info": {}}

    A._write_snapshot(snap(10_000_001, '{"new": 1}'))
    A._write_snapshot(snap(10_000_000, '{"old": 1}'))       # finished late, e.g. a background save
    assert "new" in (tmp_path / A.ANALYSIS_NAME).read_text()


def test_factor_changed_just_before_next_is_used(window, monkeypatch):
    """Phase 5 review: a factor changed less than 0.5 s before 'Next' was lost."""
    w = window
    p = w.panel
    monkeypatch.setattr(w, "message", lambda *a, **k: None)
    w.load_example()
    p.go(2)
    _wait(w)
    n1 = int(w.state.cell.sum())
    p.cell_factor.setValue(0.6)                 # starts the 0.5 s timer
    p.go(3)                                     # at once: the cell is computed first, then step 4 opens
    t = time.time()
    while (p.step != 3 or w.state.computing or w.state.result is None) and time.time() - t < 60:
        _pump(0.05)
    assert p.step == 3 and int(w.state.cell.sum()) > n1
    assert w.state.result is not None and np.array_equal(w.state.result.masks["cell"], w.state.cell)
    before = w.state.result
    p.factor["nucleus"].setValue(0.6)
    p.go(4)
    t = time.time()
    while (p.step != 4 or w.state.computing) and time.time() - t < 60:
        _pump(0.05)
    assert p.step == 4 and w.state.result is not before
    assert w.state.result_key == w._result_key()


def test_new_analysis_closes_the_image_and_quit_asks(window, monkeypatch):
    w = window
    p = w.panel
    msgs = []
    monkeypatch.setattr(w, "message", lambda text, *a, **k: msgs.append(text))
    w.load_example()
    p.go(2)
    _wait(w)
    p.go(3)
    _wait(w)
    p.go(4)
    w.accept_cell()
    assert w._unsaved()
    # quitting with unsaved cells asks; cancelling keeps the window
    from qtpy.QtWidgets import QMessageBox
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: setattr(self, "_clicked",
                        next(b for b in self.buttons() if b.text() == "Cancel")))
    monkeypatch.setattr(QMessageBox, "clickedButton", lambda self: self._clicked)
    assert not w.confirm_close()
    monkeypatch.setattr(w, "ask", lambda *a, **k: True)
    w.new_analysis()
    assert w.state.full is None and p.step == 0 and not p.next_btn.isEnabled()
    assert not w.viewer.layers
    w.accept_cell()                              # nothing open: a clear message, nothing saved
    assert msgs and "No image" in msgs[-1]
    assert not any(im.cells for im in w.session.images)


def test_start_page_and_home(window):
    """Start page first (D59); New analysis opens the image window, Home brings the start page back."""
    w = window
    w.show_start()
    _pump()
    assert w.start.isVisible() and not w.viewer.window._qt_window.isVisible()
    assert not w.start.continue_btn.isVisible()
    w.start.card_new.clicked.emit()
    _pump()
    assert w.viewer.window._qt_window.isVisible() and not w.start.isVisible()
    w.load_example()
    w.show_start()
    _pump()
    assert w.start.continue_btn.isVisible()              # an analysis is open: continue it
    w.start.continue_btn.click()
    _pump()
    assert w.viewer.window._qt_window.isVisible()
    from localization_tool.gui.start import AboutDialog, about_text

    assert "npape2@uni-muenster.de" in about_text() and "publish" in about_text()
    AboutDialog(w.start).close()


def test_step_bar_navigation(window):
    """Loading stays in step 1; the step bar opens finished steps and reachable later ones (D62)."""
    w, p = window, window.panel
    w.load_example()
    assert p.step == 0
    bar = p.header.buttons
    assert bar[0].property("step") == "current"
    assert bar[1].isEnabled() and bar[2].isEnabled()     # roles preset: channels done, cell can be opened
    assert not bar[3].isEnabled() and "step 3" in p.header.toolTip()
    bar[2].click()
    _wait(w)
    assert p.step == 2 and w.state.cell is not None
    p.refresh()
    assert bar[3].isEnabled()
    bar[0].click()                                        # back to step 1, nothing is lost
    assert p.step == 0 and w.state.cell is not None
    bar[3].click()                                        # forward over two steps: all complete
    _wait(w)
    assert p.step == 3


def test_blind_review_hides_names(window, tmp_path):
    from localization_tool.core.reader import write_imagej_tiff
    from localization_tool.core.synthetic import make_synthetic_cell

    w, p = window, window.panel
    w.session.blind = True
    path = tmp_path / "WT_treated_cell.tif"
    write_imagej_tiff(make_synthetic_cell(), path)
    w.add_images([path], open_first=True)
    t = time.time()
    while w.state.full is None and time.time() - t < 30:
        _pump(0.1)
    im = w.session.images[-1]
    row = w.session.blind_order().index(len(w.session.images) - 1)
    assert im.code and im.code in p.image_list.item(row).text() and path.name not in p.image_list.item(row).text()
    assert w.image_index(row) == len(w.session.images) - 1
    texts = [p.info.itemAt(r, QFormLayout.ItemRole.FieldRole).widget().text() for r in range(p.info.rowCount())]
    assert not any(path.name in t for t in texts) and path.name not in w.viewer.window._qt_window.windowTitle()
    assert p.reveal_btn.isVisible() and not p.meta_box.isVisible()
    w.reveal_current()
    assert im.revealed and path.name in p.image_list.item(row).text()


def test_methods_dialog_and_word_file(window, tmp_path):
    from localization_tool import __version__
    from localization_tool.core.methods_text import save_docx
    from localization_tool.gui.start import MethodsDialog

    dlg = MethodsDialog(window.viewer.window._qt_window, window.params, None, "the default settings")
    assert __version__ in dlg.text.toPlainText()
    dlg.lang.setCurrentIndex(1)
    assert "Segmentierung" in dlg.text.toPlainText()
    out = save_docx(dlg.current_text(), tmp_path / "m.docx")
    assert out.exists() and out.stat().st_size > 5000
    dlg.close()


def test_offset_table_keeps_fluorophore_and_adds_rows(window):
    """Offsets measured without an open image are stored per fluorophore (review of 1.2)."""
    from localization_tool.gui.settings import SettingsDialog

    dlg = SettingsDialog(None, window.params, [], window.session.vocabulary, program_defaults=True)
    assert dlg.off.rowCount() == 0
    names, how = dlg.bead_labels(["Alexa Fluor 647", "DAPI"])
    assert names == ["Alexa Fluor 647", "DAPI"] and "no image" in how
    dlg._off_fluor["Alexa Fluor 647"] = "Alexa Fluor 647"
    r = dlg._offset_row("Alexa Fluor 647")
    dlg.off.cellWidget(r, 1).setValue(0.13)
    params, _ = dlg.values()
    assert params.channel_offsets == [{"channel": "Alexa Fluor 647", "dz_um": 0.13, "dy_um": 0.0, "dx_um": 0.0,
                                       "fluor": "Alexa Fluor 647"}]
    dlg2 = SettingsDialog(None, params, ["TTN"], window.session.vocabulary,
                          channels=[("TTN", "Alexa Fluor 647", "Target"), ("T-PMT", "", "Ignore")])
    assert dlg2.bead_labels(["Alexa Fluor 647", "T-PMT"])[0] == ["TTN", "T-PMT"]   # by position (no fluor on T-PMT)
    assert dlg2.values()[0].channel_offsets == params.channel_offsets              # stored entries survive


def test_practice_cell_stays_out_of_real_results(window):
    """UI review of 1.2: the synthetic practice cell is exported only when nothing else is in the analysis."""
    from pathlib import Path

    from localization_tool.core import project_file
    from localization_tool.gui.app import EXAMPLE_PATH

    w = window
    row = {"image": "x", "target": "TTN", "enrichment_nucleus": 1.0}
    ex = project_file.ImageSpec(path=Path(EXAMPLE_PATH))
    ex.cells.append(project_file.CellSpec(name="cell1", status="accepted", results={"rows": [dict(row)]}))
    w.session.images = [ex]
    assert len(w.all_rows()[0]) == 1                    # practice alone: the workbook shows it
    real = project_file.ImageSpec(path=Path("/data/WT_cell.czi"))
    real.cells.append(project_file.CellSpec(name="cell1", status="accepted", results={"rows": [dict(row)]}))
    w.session.images = [ex, real]
    rows = w.all_rows()[0]
    assert len(rows) == 1 and rows[0]["status"] == "accepted"
    w.new_analysis(confirm=False)


def test_cell_viewer_pictures_values_batch_and_3d(window, tmp_path, monkeypatch):
    """D71: open a saved analysis, show plane, section and 3D, save PNG and TIFF, export all cells."""
    from qtpy.QtWidgets import QFileDialog

    from localization_tool.gui import cellviewer
    from test_v13 import make_analysis

    path = make_analysis(tmp_path / "analysis")
    v = cellviewer.CellViewer(path)
    v.show()
    assert v.cell_list.count() == 1
    v.cell_list.setCurrentRow(0)
    _pump(0.5)
    assert v.cell is not None and v.preview.pixmap() is not None and not v.preview.pixmap().isNull()
    out = tmp_path / "out"
    out.mkdir()
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(out / "pic.png"), "")))
    v.export_png()
    assert (out / "pic.png").exists()
    v.set_mode("section")
    v.angle.setValue(45)
    _pump(0.3)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(out / "sec.tif"), "")))
    v.export_tiff()
    assert (out / "sec.tif").exists()
    # export all cells into one folder per condition, with the settings used
    monkeypatch.setattr(cellviewer._BatchDialog, "exec_", lambda self: True)
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", staticmethod(lambda *a, **k: str(out)))
    monkeypatch.setattr(cellviewer.QMessageBox, "information", staticmethod(lambda *a, **k: None))
    v.export_all()
    assert list((out / "WT").glob("*.png")) and (out / "figure_settings.json").exists()
    assert (out / "exported_cells.csv").exists()
    # 3D: volumes and the cell surface, saved as shown
    v.set_mode("3d")
    _pump(1.0)
    assert v._model3d is not None and len(v._model3d.layers) >= 4
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(out / "3d.png"), "")))
    v.export_png()
    assert (out / "3d.png").exists()
    # feedback 13.10.: the saved 3D picture shows the same frame as the window, only sharper
    v.settings.scale = 1
    one = v.screenshot_3d()
    v.settings.scale = 2
    two = v.screenshot_3d()
    assert two.shape[0] == pytest.approx(2 * one.shape[0], abs=2)
    frac = [float((im.max(-1) > 5).mean()) for im in (one, two)]
    assert frac[1] == pytest.approx(frac[0], abs=0.05)
    # light from the camera: the cell surface is about as bright from below as from above
    surf = next(iter(v._surface_layers.values()))
    v._model3d.camera.angles = (0, 0, 90)
    _pump(0.2)
    sf = v._qt3d.canvas.layer_to_visual[surf].node.shading_filter
    top = tuple(sf.light_dir)
    v._model3d.camera.angles = (0, 90, 90)
    _pump(0.2)
    view = -np.array(v._model3d.camera.view_direction)[::-1]
    assert np.allclose(sf.light_dir, view, atol=1e-6) and tuple(sf.light_dir) != top
    assert sf.ambient_light.rgba[3] == pytest.approx(0.45, abs=0.01)
    v.close()
    # the viewer remembers its settings for this analysis
    v2 = cellviewer.CellViewer(path)
    assert v2.settings.mode in ("plane", "section")
    v2.close()


def test_cell_viewer_mask_panel_order_and_colours(window, tmp_path):
    """Feedback 13.10.: target mask first, zones last after a line, a colour for every mask."""
    from localization_tool.core import figure
    from localization_tool.gui import cellviewer
    from test_v13 import make_analysis

    v = cellviewer.CellViewer(make_analysis(tmp_path / "analysis"))
    v.show()
    v.cell_list.setCurrentRow(0)
    _pump(0.5)
    names = [m.name for m in v.settings.masks]
    assert names[0].startswith("target_")
    zones = [i for i, n in enumerate(names) if n in figure.PANEL_ZONES]
    assert zones and zones == list(range(zones[0], len(names)))
    cell_row = names.index("cell")
    combos = [w for w in v.findChildren(cellviewer.QComboBox) if w.findData("#FFFFFF") >= 0 and w.isVisible()]
    assert len(combos) >= len(names)
    mv = v.settings.masks[cell_row]
    v._mask_color(mv, "#FF3B30")
    _pump(0.2)
    assert mv.color == "#FF3B30"
    v.close()
