"""Command line, analysis file and batch run on synthetic data."""

import csv
import json

import numpy as np

from localization_tool.core import project_file
from localization_tool.core.channels import ChannelAssignment
from localization_tool.core.cli import assign_by_rules, main, parse_roles
from localization_tool.core.crop import CropBox
from localization_tool.core.reader import read_image, write_imagej_tiff
from localization_tool.core.synthetic import make_golden_cell


def _write_golden(tmp_path):
    stack, truth = make_golden_cell(voxel_um=(0.13, 0.08, 0.08), margin_um=1.5)
    path = tmp_path / "images" / "golden.tif"
    path.parent.mkdir()
    write_imagej_tiff(stack, path)
    return path, stack, truth


def test_parse_roles_and_assign():
    rules = parse_roles("DAPI=Nucleus,488=Actin,647=Target:TTN")
    assert rules[2] == ("647", "Target", "TTN")
    stack, _ = make_golden_cell(voxel_um=(0.13, 0.1, 0.1))
    a = assign_by_rules(stack.channels, rules)
    assert [x.role for x in a] == ["Target", "Actin", "Nucleus"]
    assert assign_by_rules(stack.channels, rules[:1]) is None


def test_analysis_file_round_trip(tmp_path):
    af = project_file.AnalysisFile(
        images=[project_file.ImageSpec(path=tmp_path / "a" / "x.czi", experiment="E1", condition="treated",
                                       cells=[project_file.CellSpec("c1", CropBox(0, 5, 1, 9, 2, 8),
                                                                    exclusions=[[(0, 0), (0, 3), (3, 3)]])])],
        panel=[ChannelAssignment("Target", "TTN", "#FF00FF")])
    path = project_file.save(af, tmp_path / "run")
    assert path.name.endswith(".lta.json")
    raw = json.loads(path.read_text())
    assert raw["images"][0]["path"] == "a/x.czi"
    back = project_file.load(path)
    assert back.images[0].path == tmp_path / "a" / "x.czi"
    assert back.images[0].cells[0].box == CropBox(0, 5, 1, 9, 2, 8)
    assert back.panel[0].label == "TTN"
    assert back.params == af.params


def test_init_and_analyze_end_to_end(tmp_path):
    img, stack, truth = _write_golden(tmp_path)
    lta = tmp_path / "golden.lta.json"
    assert main(["init", str(img.parent), "--out", str(lta), "--roles", "647=Target:TTN,488=Actin,DAPI=Nucleus"]) == 0
    af = project_file.load(lta)
    assert len(af.images) == 1 and af.images[0].assignments[0].label == "TTN"
    out = tmp_path / "results"
    assert main(["analyze", str(lta), "--out", str(out), "--csv"]) == 0
    rows = list(csv.DictReader((out / "cells.csv").open(encoding="utf-8-sig")))
    assert len(rows) == 1
    r = rows[0]
    assert r["target"] == "TTN"
    total = sum(float(r[f"signal_{k}_pct"]) for k in ("nucleus", "cytoplasm"))
    assert abs(total - 100) < 1e-3
    assert abs(float(r["signal_nucleus_pct"]) - truth["expected"]["signal_nucleus_pct"]) < 2.0
    assert (out / "profiles.csv").exists()
    pngs = list((out / "Control images").glob("*.png"))
    assert len(pngs) == 1

    # the workbook: read me first, summary, prism sheets, all values, info
    from openpyxl import load_workbook
    wb = load_workbook(out / "Results.xlsx")
    names = wb.sheetnames
    assert names[0] == "Read me" and names[1] == "Summary"
    assert "TTN Enr nucleus" in names and "All values" in names and names[-1] == "Info"
    ws = wb["TTN Enr nucleus"]
    assert ws.cell(row=2, column=1).value == "no condition"
    assert abs(ws.cell(row=3, column=1).value - float(r["enrichment_nucleus"])) < 1e-3
    summary = wb["Summary"]
    groups = [c.value for c in summary[1]]
    head = [c.value for c in summary[2]]
    assert "Nucleus" in groups and "Cytoplasm" in groups and "Enrichment" in head


def test_corrected_masks_are_used_by_the_batch(tmp_path):
    import numpy as np

    img, stack, truth = _write_golden(tmp_path)
    labels = truth["labels"].copy()
    labels[labels == 2] = 3                      # painted: no actin at all
    project_file.save_masks(tmp_path / "masks" / "c1.npz", labels=labels)
    af = project_file.AnalysisFile(
        images=[project_file.ImageSpec(path=img, cells=[project_file.CellSpec("c1", mask_file="masks/c1.npz")],
                                       assignments=[ChannelAssignment("Target", "TTN", "#FF00FF"),
                                                    ChannelAssignment("Actin", "Actin", "#00FF00"),
                                                    ChannelAssignment("Nucleus", "Nucleus", "#0000FF")])])
    lta = project_file.save(af, tmp_path / "corr")
    from localization_tool.core.batch import run
    res = run(project_file.load(lta), tmp_path / "out", control_images=False, csv_tables=True, base_dir=tmp_path)
    rows = list(csv.DictReader(res["cells_csv"].open(encoding="utf-8-sig")))
    assert rows[0]["volume_actin_pct"] == ""             # no actin: no value, not 0 % (phase 5 review)
    assert "no actin found" in rows[0]["notes"]
    assert np.isclose(float(rows[0]["volume_nucleus_pct"]), 100 * truth["nucleus"].sum() / truth["cell"].sum(), atol=0.01)


def test_analyze_reports_missing_file(tmp_path):
    af = project_file.AnalysisFile(images=[project_file.ImageSpec(path=tmp_path / "missing.czi")],
                                   panel=[ChannelAssignment("Target", "T", "#FFFFFF")])
    lta = project_file.save(af, tmp_path / "m")
    assert main(["analyze", str(lta), "--out", str(tmp_path / "o"), "--no-images"]) == 1


def test_tiff_round_trip_keeps_values(tmp_path):
    img, stack, _ = _write_golden(tmp_path)
    back = read_image(img)
    assert np.array_equal(back.data, stack.data)
