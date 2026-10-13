"""Version 2.1: comparison channels (D83), polarity confidence (D84), rear to front (D85), envelope halves
and peak (D86), polarity axis from the actin-rich mask (D88), target values from the target mask (D89),
crash log (D87). Synthetic cells with known geometry."""

import json

import numpy as np
import pytest

from localization_tool.core import batch, distances, metrics, polarity, project_file
from localization_tool.core.analysis import analyze_cell
from localization_tool.core.channels import ChannelAssignment
from localization_tool.core.image import ChannelInfo, ImageStack
from localization_tool.core.params import AnalysisParams
from localization_tool.core.synthetic import make_golden_cell

ASSIGN = [ChannelAssignment("Target", "T", "#FF2BD6"), ChannelAssignment("Actin", "Actin", "#2BD46B"),
          ChannelAssignment("Nucleus", "Nucleus", "#3D7DFF")]
ALL = AnalysisParams(target_signal="all", sensitivity=False)
MASK = AnalysisParams(target_signal="mask", sensitivity=False)


def _with_channel(stack, index, values, extra=False):
    data = stack.data.astype(np.float32)
    values = np.asarray(values, np.float32)
    if extra:
        data = np.concatenate([data, values[None]], axis=0)
        chans = list(stack.channels) + [ChannelInfo(index=len(stack.channels), name="VIM")]
    else:
        data[index] = values
        chans = list(stack.channels)
    return ImageStack(np.clip(data, 0, 65535).astype(np.uint16), stack.voxel_um, chans, info=dict(stack.info))


def _true(truth):
    return {k: truth[k] for k in ("cell", "nucleus", "actin")}


# ---------------------------------------------------------------------------- D89
def test_mask_counting_ignores_diffuse_background():
    """A bright block in the cytoplasm on an even haze: all signal counts the haze in the nucleus, the
    target mask leaves it out, so the nucleus share falls to about 0."""
    stack, truth = make_golden_cell()
    cell, nuc = truth["cell"], truth["nucleus"]
    t = np.where(cell, 20.0, 0.0)
    block = cell & ~nuc & (np.arange(stack.data.shape[3])[None, None, :] > stack.data.shape[3] * 0.75)
    t[block] = 2000.0
    s = _with_channel(stack, 0, t)
    a = analyze_cell(s, ASSIGN, ALL, masks=_true(truth)).target_values["T"]
    m = analyze_cell(s, ASSIGN, MASK, masks=_true(truth)).target_values["T"]
    assert a["signal_nucleus_pct"] > 1.0
    assert m["signal_nucleus_pct"] < 0.1 * a["signal_nucleus_pct"]
    assert m["signal_nucleus_pct"] + m["signal_cytoplasm_pct"] == pytest.approx(100.0, abs=1e-6)
    # the volume shares stay those of the compartments
    assert m["volume_nucleus_pct"] == pytest.approx(a["volume_nucleus_pct"], abs=1e-9)


def test_mask_counting_equals_all_signal_without_background():
    stack, truth = make_golden_cell()
    cell, nuc = truth["cell"], truth["nucleus"]
    t = np.zeros(cell.shape, np.float32)
    t[nuc] = 800.0
    s = _with_channel(stack, 0, t)
    a = analyze_cell(s, ASSIGN, ALL, masks=_true(truth)).target_values["T"]
    m = analyze_cell(s, ASSIGN, MASK, masks=_true(truth)).target_values["T"]
    assert m["signal_nucleus_pct"] == pytest.approx(a["signal_nucleus_pct"], abs=0.5)


def test_unstained_target_counted_in_its_mask_has_no_values():
    stack, truth = make_golden_cell()
    noise = np.random.default_rng(1).poisson(3.0, size=stack.data[0].shape)
    r = analyze_cell(_with_channel(stack, 0, noise), ASSIGN, MASK)
    v = r.target_values["T"]
    assert not r.masks["target_T"].any()
    assert np.isnan(v["signal_nucleus_pct"]) and np.isnan(v["enrichment_cytoplasm"])
    assert any("no staining above noise" in n for n in r.notes)


def test_target_edge_enters_the_sensitivity_span_only_when_counted_in_the_mask():
    stack, truth = make_golden_cell()
    p_mask = AnalysisParams(target_signal="mask")
    p_all = AnalysisParams(target_signal="all")
    m = analyze_cell(stack, ASSIGN, p_mask, masks=_true(truth))
    a = analyze_cell(stack, ASSIGN, p_all, masks=_true(truth))
    vm, va = m.target_values["T"], a.target_values["T"]
    assert vm["signal_nucleus_pct_min"] <= vm["signal_nucleus_pct"] <= vm["signal_nucleus_pct_max"]
    assert vm["signal_nucleus_pct_max"] - vm["signal_nucleus_pct_min"] > 0       # the target edge moves
    assert np.isnan(va["signal_nucleus_pct_min"])                                 # given masks, no variant
    assert m.cell_values["sensitivity_variants"] == 2


def test_older_analysis_keeps_all_signal_and_the_old_axis(tmp_path):
    af = project_file.AnalysisFile(params=AnalysisParams())
    path = project_file.save(af, tmp_path / "a")
    raw = json.loads(path.read_text())
    for key in ("target_signal", "polarity_from", "polarity_clear_index", "envelope_lateral_only", "comparison"):
        raw["params"].pop(key)
    path.write_text(json.dumps(raw))
    back = project_file.load(path)
    assert back.params.target_signal == "all" and back.params.polarity_from == "actin_signal"
    assert AnalysisParams().target_signal == "mask" and AnalysisParams().polarity_from == "actin_rich"
    # a cell stored with the old settings and the current result version needs no new computation
    stored = back.params.to_dict()
    for key in ("target_signal", "polarity_from", "polarity_clear_index", "envelope_lateral_only", "comparison"):
        stored.pop(key)
    spec = project_file.CellSpec("c1", results={"rows": [{}], "params": stored,
                                                "results_version": batch.RESULTS_VERSION})
    assert not batch.needs_update(spec, back.params)


# ---------------------------------------------------------------------------- D86
@pytest.mark.parametrize("d0", [0.3, -0.3])
def test_envelope_peak_finds_the_side_of_a_thin_shell(d0):
    stack, truth = make_golden_cell()
    d = distances.signed_distance(truth["nucleus"], stack.voxel_um)
    t = np.where(truth["cell"], 3000.0 * np.exp(-0.5 * ((d - d0) / 0.08) ** 2), 0.0)
    v = analyze_cell(_with_channel(stack, 0, t), ASSIGN, ALL, masks=_true(truth)).target_values["T"]
    assert v["envelope_peak_um"] == pytest.approx(d0, abs=0.1)
    inner, outer = v["enrichment_envelope_inner"], v["enrichment_envelope_outer"]
    assert (outer > inner) if d0 > 0 else (inner > outer)


def test_envelope_halves_add_up_to_the_zone():
    stack, truth = make_golden_cell()
    v = analyze_cell(stack, ASSIGN, ALL, masks=_true(truth)).target_values["T"]
    for q in ("signal", "volume"):
        assert v[f"{q}_envelope_inner_pct"] + v[f"{q}_envelope_outer_pct"] == pytest.approx(
            v[f"{q}_envelope_zone_pct"], abs=1e-6)


def test_lateral_envelope_keeps_the_side_of_a_sphere():
    """On a sphere the side within 45° of the equator holds sin(45°) = 71 % of a thin shell."""
    stack, truth = make_golden_cell()
    nuc = truth["nucleus"]
    d = distances.signed_distance(nuc, stack.voxel_um)
    shell = np.abs(d) <= 0.15
    lat = metrics.lateral_mask(d, stack.voxel_um)
    assert (shell & lat).sum() / shell.sum() == pytest.approx(np.sin(np.radians(45)), abs=0.08)
    p = AnalysisParams(target_signal="all", sensitivity=False, envelope_lateral_only=True)
    on = analyze_cell(stack, ASSIGN, p, masks=_true(truth)).target_values["T"]
    off = analyze_cell(stack, ASSIGN, ALL, masks=_true(truth)).target_values["T"]
    assert on["volume_envelope_zone_pct"] < 0.85 * off["volume_envelope_zone_pct"]
    assert on["signal_nucleus_pct"] == pytest.approx(off["signal_nucleus_pct"], abs=1e-9)   # only the zones


def test_profile_peak_parabola():
    rows = [{"bin_start_um": a, "bin_end_um": a + 0.2, "enrichment": e}
            for a, e in zip((-0.4, -0.2, 0.0, 0.2, 0.4), (1.0, 2.0, 4.0, 3.0, 1.0))]
    pk = metrics.profile_peak(rows)
    assert 0.1 < pk < 0.2                    # between the best bin (0.1) and its higher neighbour (0.3)
    assert np.isnan(metrics.profile_peak(rows[:2]))
    rising = [{"bin_start_um": a, "bin_end_um": a + 0.2, "enrichment": 1.0 + a}
              for a in np.arange(-1.2, 1.2, 0.2)]
    assert np.isnan(metrics.profile_peak(rising))        # no maximum at the edge, only a slope


# ---------------------------------------------------------------------------- D83
def test_comparison_channel_area_and_target_inside_it():
    stack, truth = make_golden_cell()
    cell, nuc = truth["cell"], truth["nucleus"]
    nx = stack.data.shape[3]
    area = cell & ~nuc & (np.arange(nx)[None, None, :] > nx * 0.7)
    vim = np.where(area, 1500.0, 0.0)
    t = np.where(area, 900.0, 0.0)
    s = _with_channel(_with_channel(stack, 0, t), None, vim, extra=True)
    assign = ASSIGN + [ChannelAssignment("Comparison", "VIM", "#FFD60A")]
    r = analyze_cell(s, assign, MASK, masks=_true(truth))
    v = r.target_values["T"]
    phi = 100.0 * area.sum() / cell.sum()
    # the smoothed mask edge (σ 0.15 µm) widens the area a little
    assert v["volume_VIM_area_pct"] == pytest.approx(phi, rel=0.15)
    assert v["signal_in_VIM_pct"] == pytest.approx(100.0, abs=1.0)
    assert v["enrichment_in_VIM"] == pytest.approx(v["signal_in_VIM_pct"] / v["volume_VIM_area_pct"], rel=1e-9)
    assert v["enrichment_in_VIM"] == pytest.approx(100.0 / phi, rel=0.15)
    assert v["target_mask_in_VIM_pct"] > 95
    assert v["role"] == "Target" and r.target_values["VIM"]["role"] == "Comparison"
    assert "pcc_VIM" in v                                     # colocalization as between two targets
    assert "signal_in_VIM_pct" not in r.target_values["VIM"]  # no comparison with itself
    assert r.masks["target_VIM"].any()


def test_cells_without_comparison_channel_keep_their_values():
    stack, truth = make_golden_cell()
    a = analyze_cell(stack, ASSIGN, MASK, masks=_true(truth)).target_values["T"]
    assert not any(k.startswith("signal_in_") for k in a)


# ---------------------------------------------------------------------------- D84, D85, D88
def _cap_cell():
    stack, truth = make_golden_cell()
    return stack, truth


def test_rear_to_front_ratio_of_an_even_and_a_rear_target():
    stack, truth = _cap_cell()
    r = analyze_cell(stack, ASSIGN, ALL, masks=_true(truth))
    axis = None
    cell = truth["cell"]
    pol = polarity.polarity(cell, stack.data[1].astype(float), truth["nucleus"], stack.voxel_um)
    axis = pol["_axis"]
    even = polarity.target_front(np.where(cell, 1.0, 0.0), cell, axis, stack.voxel_um)
    assert even["rear_to_front_ratio"] == pytest.approx(1.0, abs=1e-6)
    front = polarity.front_half(cell, axis["centre_um"], axis["unit"], stack.voxel_um)
    rear_only = polarity.target_front(np.where(cell & ~front, 1.0, 0.0), cell, axis, stack.voxel_um)
    assert np.isinf(rear_only["rear_to_front_ratio"]) or np.isnan(rear_only["rear_to_front_ratio"]) \
        or rear_only["rear_to_front_ratio"] > 10
    mixed = polarity.target_front(np.where(cell & ~front, 3.0, np.where(cell, 1.0, 0.0)), cell, axis, stack.voxel_um)
    assert mixed["rear_to_front_ratio"] == pytest.approx(3.0, rel=1e-6)
    assert "rear_to_front_ratio" in r.target_values["T"]


def test_axis_from_the_actin_rich_mask_follows_the_mask():
    """D88: the axis points to the actin-rich mask; a different mask turns the axis."""
    stack, truth = _cap_cell()
    cell = truth["cell"]
    actin = stack.data[1].astype(float)
    nx = cell.shape[2]
    right = cell & (np.arange(nx)[None, None, :] > nx * 0.7)
    left = cell & (np.arange(nx)[None, None, :] < nx * 0.3)
    pr = polarity.polarity(cell, actin + 1.0, truth["nucleus"], stack.voxel_um, front_mask=right)
    pl = polarity.polarity(cell, actin + 1.0, truth["nucleus"], stack.voxel_um, front_mask=left)
    assert abs(pr["polarity_axis_deg"]) < 20
    assert abs(abs(pl["polarity_axis_deg"]) - 180) < 20
    assert pr["polarity_index"] > 0.3


def test_polarity_clear_flag_and_arrow_length():
    stack, truth = _cap_cell()
    low = analyze_cell(stack, ASSIGN, AnalysisParams(target_signal="all", sensitivity=False,
                                                     polarity_clear_index=5.0), masks=_true(truth))
    high = analyze_cell(stack, ASSIGN, AnalysisParams(target_signal="all", sensitivity=False,
                                                      polarity_clear_index=0.01), masks=_true(truth))
    assert low.cell_values["polarity_clear"] == 0.0 and high.cell_values["polarity_clear"] == 1.0
    a, b = low.overlays["polarity"], high.overlays["polarity"]
    length = lambda o: float(np.hypot(*(np.array(o["end"]) - np.array(o["start"]))))  # noqa: E731
    assert length(a) < length(b)
    assert not a["clear"] and b["clear"] and a["width"] < b["width"]


# ---------------------------------------------------------------------------- export and texts
def test_workbook_has_comparison_columns_and_sheet(tmp_path):
    from openpyxl import load_workbook

    from localization_tool.core.crop import CropBox
    from localization_tool.core.export_excel import write_results

    stack, truth = make_golden_cell()
    nx = stack.data.shape[3]
    area = truth["cell"] & (np.arange(nx)[None, None, :] > nx * 0.6)
    s = _with_channel(stack, None, np.where(area, 1500.0, 0.0), extra=True)
    assign = ASSIGN + [ChannelAssignment("Comparison", "VIM", "#FFD60A")]
    r = analyze_cell(s, assign, MASK, masks=_true(truth))
    rows, prof = batch.result_rows("x.czi", "E1", "WT", "cell1", "accepted", CropBox(0, 1, 0, 1, 0, 1), r)
    path = write_results(tmp_path / "Results.xlsx", rows, prof, {"params": MASK.to_dict()})
    wb = load_workbook(path)
    assert "T Enr in VIM" in wb.sheetnames
    heads = [c.value for c in wb["Summary"][1]]
    assert "In VIM area" in heads and "Rear half" in heads and "Envelope zone, outer half" in heads


def test_methods_text_names_the_counting_and_the_comparison():
    from localization_tool.core.methods_text import Summary, methods_text

    s = Summary(roles={"T": "Target", "VIM": "Comparison"}, targets=["T"])
    en = methods_text(AnalysisParams(), s, "en")
    de = methods_text(AnalysisParams(), s, "de")
    assert "only the signal inside this mask" in en and "nur das Signal in dieser Maske" in de
    assert "comparison channel" in en and "Vergleichskanal" in de
    assert "inner and an outer half" in en and "rear half" in en
    legacy = methods_text(AnalysisParams(target_signal="all", polarity_from="actin_signal"), s, "en")
    assert "the main values do not depend on it" in legacy and "actin-weighted centroid of the cell" in legacy


# ---------------------------------------------------------------------------- D87 crash log
def test_previous_crash_text_recognises_native_crashes_and_missing_exit():
    from localization_tool.gui.app import CLEAN_EXIT, previous_crash_text

    start = "CellLocalizer 2.1.0 started 2026-10-13 09:00:00\n"
    assert previous_crash_text(start + f"{CLEAN_EXIT} 2026-10-13 10:00:00\n") is None
    assert previous_crash_text(start + "Windows fatal exception: access violation\n")
    assert "no clean exit" in previous_crash_text(start)
    assert previous_crash_text("") is None
