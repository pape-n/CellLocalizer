"""Tests for the findings of the phase 5 review (version 1.0.0): wrong labels, missing values,
damaged files, quality notes, Excel layout and robustness."""

import numpy as np
import pytest

from localization_tool.core import batch, project_file
from localization_tool.core.analysis import analyze_cell, cell_mask
from localization_tool.core.channels import ChannelAssignment, suggest_assignments
from localization_tool.core.image import ChannelInfo, ImageStack
from localization_tool.core.reader import ReaderError, read_image, write_imagej_tiff
from localization_tool.core.synthetic import make_golden_cell

ASSIGN = [ChannelAssignment("Target", "T", "#FF2BD6"), ChannelAssignment("Actin", "Actin", "#2BD46B"),
          ChannelAssignment("Nucleus", "Nucleus", "#3D7DFF")]


def _noise_target(stack):
    data = stack.data.copy()
    data[0] = np.random.default_rng(1).poisson(3.0, size=data[0].shape).astype(data.dtype)
    return ImageStack(data, stack.voxel_um, stack.channels, info=dict(stack.info))


def test_cell_thresholds_keep_their_channel_names():
    """A channel without staining is not used for the cell edge; the other thresholds must still carry
    the right names (before: shifted by one channel)."""
    stack, _ = make_golden_cell()
    _, thr_ref = cell_mask(stack, ASSIGN)
    _, thr = cell_mask(_noise_target(stack), ASSIGN)
    assert np.isnan(thr["cell_T"])
    assert thr["cell_Nucleus"] == pytest.approx(thr_ref["cell_Nucleus"], rel=0.05)
    assert thr["cell_Actin"] == pytest.approx(thr_ref["cell_Actin"], rel=0.05)


def test_missing_nucleus_channel_gives_no_value_not_zero():
    stack, _ = make_golden_cell()
    roles = [ASSIGN[0], ASSIGN[1], ChannelAssignment("Ignore", "Nucleus", "#3D7DFF")]
    r = analyze_cell(stack, roles)
    v = r.target_values["T"]
    for k in ("signal_nucleus_pct", "signal_cytoplasm_pct", "enrichment_nucleus", "signal_envelope_zone_pct",
              "signal_cytoplasm_without_actin_pct"):
        assert np.isnan(v[k]), k
    assert np.isnan(r.cell_values["nucleus_volume_um3"])
    assert v["signal_actin_pct"] > 0                       # actin is still measured
    assert any("no nucleus channel" in n for n in r.notes)


def test_missing_actin_channel_gives_no_value_not_zero():
    stack, _ = make_golden_cell()
    roles = [ASSIGN[0], ChannelAssignment("Ignore", "Actin", "#2BD46B"), ASSIGN[2]]
    r = analyze_cell(stack, roles)
    v = r.target_values["T"]
    for k in ("signal_actin_pct", "enrichment_actin", "signal_actin_rich_pct", "signal_cytoplasm_without_actin_pct"):
        assert np.isnan(v[k]), k
    assert v["signal_nucleus_pct"] + v["signal_cytoplasm_pct"] == pytest.approx(100.0)
    assert any("no actin channel" in n for n in r.notes)


def test_notes_for_unstained_target_two_nuclei_and_missing_calibration():
    stack, truth = make_golden_cell()
    r = analyze_cell(_noise_target(stack), ASSIGN)
    assert any("T: no staining above noise" in n for n in r.notes)
    r = analyze_cell(stack, ASSIGN)
    assert not any("no staining" in n or "two cells" in n or "calibration" in n for n in r.notes)
    # a second nucleus (neighbour cell) inside the cell mask
    nuc2 = np.zeros_like(truth["nucleus"])
    zz, yy, xx = np.nonzero(truth["nucleus"])
    shift = int(round(4.4 / stack.voxel_um[2]))
    keep = xx + shift < nuc2.shape[2]
    nuc2[zz[keep], yy[keep], xx[keep] + shift] = True
    masks = {"cell": truth["cell"] | nuc2, "nucleus": truth["nucleus"] | nuc2}
    r = analyze_cell(stack, ASSIGN, masks=masks)
    assert any("possibly two cells" in n for n in r.notes)
    s2 = ImageStack(stack.data, stack.voxel_um, stack.channels, info=dict(stack.info), calibrated=False)
    assert any("no voxel calibration" in n for n in analyze_cell(s2, ASSIGN).notes)


def test_dapi_outside_the_nucleus_mask_is_reported():
    stack, truth = make_golden_cell()
    small = truth["nucleus"].copy()
    small[:, :, : int(np.flatnonzero(small.any(axis=(0, 1))).mean())] = False     # half of the nucleus
    r = analyze_cell(stack, ASSIGN, masks={"cell": truth["cell"], "nucleus": small})
    assert any("DAPI signal lies outside the nucleus mask" in n for n in r.notes)


def test_damaged_files_raise_reader_errors(tmp_path):
    czi = tmp_path / "cut.czi"
    czi.write_bytes(b"ZISRAWFILE" + b"\0" * 990)
    with pytest.raises(ReaderError):
        read_image(czi)
    stack, _ = make_golden_cell(voxel_um=(0.13, 0.1, 0.1))
    tif = tmp_path / "golden.tif"
    write_imagej_tiff(stack, tif)
    raw = tif.read_bytes()
    cut = tmp_path / "cut.tif"
    cut.write_bytes(raw[: len(raw) // 2])
    with pytest.raises(ReaderError):
        read_image(cut)


def test_one_damaged_file_does_not_stop_the_run(tmp_path):
    stack, _ = make_golden_cell(voxel_um=(0.13, 0.1, 0.1))
    good = tmp_path / "good.tif"
    write_imagej_tiff(stack, good)
    bad = tmp_path / "bad.czi"
    bad.write_bytes(b"ZISRAWFILE" + b"\0" * 990)
    af = project_file.AnalysisFile(images=[project_file.ImageSpec(path=bad), project_file.ImageSpec(path=good)],
                                   panel=ASSIGN)
    res = batch.run(af, tmp_path / "out", control_images=False)
    assert res["rows"] == 1 and len(res["errors"]) == 1 and "bad.czi" in res["errors"][0]["error"]


def test_raw_folders_are_whole_words(tmp_path):
    from localization_tool.core.files import is_raw_folder_name

    assert is_raw_folder_name("raw data") and is_raw_folder_name("Raw") and is_raw_folder_name("non processed")
    assert is_raw_folder_name("2026_unprocessed") and is_raw_folder_name("non-processed")
    assert is_raw_folder_name("RawData")
    assert not is_raw_folder_name("Neutrophil crawling assay") and not is_raw_folder_name("Withdrawn")


def test_same_file_name_in_two_folders_gets_two_control_images():
    a, b = "/d1/New-01.czi", "/d2/New-01.czi"
    na, nb = batch.control_name(a, "cell1", [a, b]), batch.control_name(b, "cell1", [a, b])
    assert na != nb and na.startswith("New-01_")
    assert batch.control_name(a, "cell1", [a]) == "New-01_cell1.png"


def test_prism_sheets_hold_checked_cells_with_ids(tmp_path):
    from openpyxl import load_workbook

    from localization_tool.core.export_excel import write_results

    def row(cell, status, value):
        return {"image": "a.czi", "experiment": "e1", "condition": "WT", "cell": cell, "target": "T",
                "status": status, "enrichment_nucleus": value, "notes": "tab\there\x0b"}

    rows = [row("cell1", "accepted", 1.5), row("cell2", "automatic", 9.0), row("cell3", "accepted, corrected", 2.5)]
    path = write_results(tmp_path / "R.xlsx", rows, [], {"params": {"envelope_um": 0.4}})
    wb = load_workbook(path)
    ws = wb["T Enr nucleus"]
    assert [ws.cell(row=i, column=1).value for i in (3, 4, 5)] == [1.5, 2.5, None]
    assert ws.cell(row=3, column=3).value == "e1 / a.czi / cell1"
    assert wb["Replicates"].cell(row=2, column=4).value == 2
    statuses = [r[5] for r in wb["Summary"].iter_rows(min_row=3, values_only=True)]
    assert "to check" in statuses and "automatic" not in statuses
    readme = " ".join(str(c.value) for r in wb["Read me"].iter_rows() for c in r if c.value)
    assert "within 0.4 µm" in readme
    assert not (tmp_path / "R.tmp.xlsx").exists()
    # no checked cell at all (command line): all cells, and the sheet says so
    path = write_results(tmp_path / "R2.xlsx", [row("cell2", "automatic", 9.0)], [], {})
    ws = load_workbook(path)["T Enr nucleus"]
    assert ws.cell(row=3, column=1).value == 9.0 and "ALL cells" in ws.cell(row=1, column=1).value


def test_analysis_file_with_bom_and_damaged_files(tmp_path):
    af = project_file.AnalysisFile(images=[project_file.ImageSpec(path=tmp_path / "x.czi")])
    p = project_file.save(af, tmp_path / "a")
    p.write_bytes(b"\xef\xbb\xbf" + p.read_bytes())
    assert project_file.load(p).images[0].path.name == "x.czi"
    for text in ('[1, 2]', '{"images": [{"experiment": "e"}]}', '{"images": [{"path": "x", "cells": [{"box": [1]}]}]}'):
        p.write_text(text, encoding="utf-8")
        with pytest.raises(ValueError):
            project_file.load(p)


def test_results_of_older_versions_are_computed_again():
    spec = project_file.CellSpec("cell1", results={"rows": [{}], "params": batch.AnalysisParams().to_dict()})
    assert batch.needs_update(spec, batch.AnalysisParams())
    spec.results["results_version"] = batch.RESULTS_VERSION
    assert not batch.needs_update(spec, batch.AnalysisParams())


def test_roles_guessed_from_the_file_name():
    chans = [ChannelInfo(0, "AF647-T8", "Alexa Fluor 647"), ChannelInfo(1, "AF555", "Alexa Fluor 555", color="#FF7E00"),
             ChannelInfo(2, "AF488-T10", "Alexa Fluor 488", color="#00FF00"), ChannelInfo(3, "DAPI-T11", "DAPI")]
    a = suggest_assignments(chans, "06_25_2026 Cells KO+vehicle+IL8  TTN647 Vim555 Act488 DAPI03")
    assert [(x.role, x.label) for x in a] == [("Target", "TTN"), ("Target", "Vim"), ("Actin", "Actin"),
                                              ("Nucleus", "Nucleus")]
    a = suggest_assignments(chans, "19_02_2026 Cells 488Act 647TTN DAPI-01-Airyscan Processing-01")
    assert [x.role for x in a] == ["Target", "Other", "Actin", "Nucleus"]
    a = suggest_assignments(chans, "WT TTN 647 DAPI-1")
    assert a[0].label == "TTN"
    for name in ("Cells fMLP 647 TTN 555 Vim 488 Act DAPI", "Cells_C5a_647-TTN_555-Vim_488-Act"):
        a = suggest_assignments(chans, name)                # wavelength first: the word after it
        assert [(x.role, x.label) for x in a][:3] == [("Target", "TTN"), ("Target", "Vim"), ("Actin", "Actin")]
    colors = [x.color for x in suggest_assignments(chans, "")]
    assert len(set(colors)) == 4                           # 647 and 555 no longer both orange
    assert [x.role for x in suggest_assignments(chans)] == ["Other", "Other", "Other", "Nucleus"]
